"""Readable Excel tables without changing the underlying review evidence."""
from __future__ import annotations

import unicodedata

from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.hyperlink import Hyperlink

FONT = "Arial"
FONT_SIZE = 11
LINE_HEIGHT = 16.5
MAX_ROW_HEIGHT = 409
OVERFLOW_SHEET = "長文の続き"
OVERFLOW_HEADERS = ["元のシート", "元のセル", "項目", "続き", "全文（分割表示）"]
HEADER_FILL = PatternFill("solid", fgColor="243B53")
ALTERNATE_FILL = PatternFill("solid", fgColor="F0F5F9")
RULE = Border(bottom=Side(style="hair", color="D9E2EC"))


def character_width(character: str) -> int:
    if unicodedata.combining(character) or character == "\r":
        return 0
    if character == "\t":
        return 4
    return 2 if unicodedata.east_asian_width(character) in ("W", "F") else 1


def display_lines(value, width: float) -> int:
    """Conservative Arial-11 line estimate, including explicit blank lines."""
    if value is None or value == "":
        return 1
    capacity = max(1, width * 0.9 - 3)
    lines, used = 1, 0
    for character in str(value):
        if character == "\n":
            lines += 1
            used = 0
        else:
            units = character_width(character)
            if used and used + units > capacity:
                lines += 1
                used = 0
            used += units
    return lines


def split_display_text(value: str, width: float, max_lines: int = 20) -> list[str]:
    """Keep every character while making chunks that fit a visible Excel row."""
    capacity = max(1, width * 0.9 - 3)
    chunks, start, lines, used = [], 0, 1, 0
    for index, character in enumerate(value):
        units = character_width(character)
        if character != "\n" and used and used + units > capacity:
            lines += 1
            used = 0
        if lines > max_lines or index - start >= 30000:
            chunks.append(value[start:index])
            start, lines, used = index, 1, 0
        if character == "\n":
            lines += 1
            used = 0
        else:
            used += units
    if start < len(value):
        chunks.append(value[start:])
    return chunks or [""]


def set_cell_value(cell, value):
    if isinstance(value, str):
        cell.value = ILLEGAL_CHARACTERS_RE.sub("", value)[:32767]
        # Review text is data, including text beginning with an equals sign.
        cell.data_type = "s"
    else:
        cell.value = value


def style_data_row(sheet, row_number: int, widths: list[float], hidden: set[int]):
    lines = 1
    for column, cell in enumerate(sheet[row_number], 1):
        cell.font = Font(name=FONT, size=FONT_SIZE, color="243B53")
        cell.alignment = Alignment(vertical="top", wrap_text=True)
        cell.border = RULE
        if isinstance(cell.value, float):
            cell.number_format = "0.00"
        if row_number % 2 == 0:
            cell.fill = ALTERNATE_FILL
        if column not in hidden:
            lines = max(lines, display_lines(cell.value, widths[column - 1]))
    sheet.row_dimensions[row_number].height = min(MAX_ROW_HEIGHT, max(30, lines * LINE_HEIGHT + 12))


def internal_link(cell, sheet_name: str, address: str):
    escaped = sheet_name.replace("'", "''")
    cell.hyperlink = Hyperlink(ref=cell.coordinate, location=f"'{escaped}'!{address}")


def append_overflow(book, source, records):
    if not records:
        return
    # A link on the entire body is underlined by spreadsheet applications.
    # Put the navigation beside the data so paragraphs stay readable.
    link_column = source.max_column + 1
    header = source.cell(1, link_column, "全文の参照")
    header.fill = HEADER_FILL
    header.font = Font(name=FONT, size=FONT_SIZE, bold=True, color="FFFFFF")
    header.alignment = Alignment(vertical="center", wrap_text=True)
    source.column_dimensions[get_column_letter(link_column)].width = 18
    if OVERFLOW_SHEET not in book.sheetnames:
        detail = write_table(book, OVERFLOW_SHEET, OVERFLOW_HEADERS, [],
                             widths=[22, 12, 26, 8, 88], freeze_panes="E2")
        detail["E1"].comment = Comment("同じ元シート・元セルの行を、続き番号の順に読むと全文になります。", "google-maps-reviews")
    detail = book[OVERFLOW_SHEET]
    for cell, title, original in records:
        start = detail.max_row + 1
        for number, chunk in enumerate(split_display_text(original, 88), 1):
            row = detail.max_row + 1
            for column, value in enumerate((source.title, cell.coordinate, title, number, chunk), 1):
                set_cell_value(detail.cell(row, column), value)
            style_data_row(detail, row, [22, 12, 26, 8, 88], set())
            internal_link(detail.cell(row, 2), source.title, cell.coordinate)
        reference = source.cell(cell.row, link_column)
        if reference.value is None:
            set_cell_value(reference, "全文を表示")
            internal_link(reference, OVERFLOW_SHEET, f"E{start}")
            reference.font = Font(name=FONT, size=FONT_SIZE, color="2563EB", underline="single")
            reference.alignment = Alignment(vertical="top", wrap_text=True)
            reference.border = RULE
            if cell.row % 2 == 0:
                reference.fill = ALTERNATE_FILL
        cell.comment = Comment(f"行高の上限を超える長文です。全文は「長文の続き」のE{start}から分割表示しています。右の「全文の参照」から移動できます。", "google-maps-reviews")
    source.auto_filter.ref = source.dimensions
    detail.auto_filter.ref = detail.dimensions
    detail.print_area = detail.dimensions
    book.move_sheet(detail, offset=len(book.worksheets) - book.worksheets.index(detail) - 1)


def write_table(book, name: str, headers, rows, *, widths=None, hidden_columns=(),
                freeze_panes="A2", print_columns=None):
    """Create every sheet directly so copied row styles cannot be lost."""
    sheet = book.create_sheet(name)
    headers = list(headers)
    widths = list(widths) if widths is not None else [24] * len(headers)
    if len(widths) != len(headers):
        raise ValueError("Excelの列幅と見出しの数が一致していません。")
    hidden = set(hidden_columns)
    sheet.sheet_view.showGridLines = False
    sheet.sheet_view.zoomScale = 90
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.sheet_properties.outlinePr.summaryRight = True
    sheet.freeze_panes = freeze_panes
    sheet.row_dimensions[1].height = max(36, max(display_lines(h, w) for h, w in zip(headers, widths)) * LINE_HEIGHT + 10)
    for column, (header, width) in enumerate(zip(headers, widths), 1):
        letter = get_column_letter(column)
        sheet.column_dimensions[letter].width = width
        sheet.column_dimensions[letter].hidden = column in hidden
        sheet.column_dimensions[letter].outlineLevel = 1 if column in hidden else 0
        cell = sheet.cell(1, column)
        set_cell_value(cell, header)
        cell.fill = HEADER_FILL
        cell.font = Font(name=FONT, size=FONT_SIZE, bold=True, color="FFFFFF")
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    overflow = []
    for row_number, values in enumerate(rows, 2):
        values = list(values)
        if len(values) != len(headers):
            raise ValueError("Excelのデータ列数と見出しの数が一致していません。")
        for column, value in enumerate(values, 1):
            cell = sheet.cell(row_number, column)
            set_cell_value(cell, value)
            if column not in hidden and isinstance(value, str):
                original = ILLEGAL_CHARACTERS_RE.sub("", value)
                if display_lines(original, widths[column - 1]) * LINE_HEIGHT + 12 > MAX_ROW_HEIGHT or len(original) > 32767:
                    overflow.append((cell, headers[column - 1], original))
        style_data_row(sheet, row_number, widths, hidden)
    sheet.auto_filter.ref = sheet.dimensions
    sheet.sheet_properties.tabColor = "243B53"
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.print_options.horizontalCentered = True
    sheet.print_title_rows = "1:1"
    first, last = print_columns or (1, len(headers))
    sheet.print_area = f"{get_column_letter(first)}1:{get_column_letter(last)}{sheet.max_row}"
    sheet.page_margins.left = sheet.page_margins.right = 0.25
    sheet.page_margins.top = sheet.page_margins.bottom = 0.4
    sheet.oddFooter.center.text = "&P / &N"
    append_overflow(book, sheet, overflow)
    return sheet
