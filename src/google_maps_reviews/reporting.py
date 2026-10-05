"""Evidence-bound review analysis, deterministic aggregation and offline reports.

Semantic annotations are written by the Skill, never inferred from star ratings.
"""
from __future__ import annotations

import hashlib
import html
import json
import re
import csv
import io
from collections import Counter
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from . import __version__

CATEGORIES = {
    "quality": "商品・料理・施術の品質", "service": "接客・スタッフ",
    "waiting": "待ち時間・提供速度", "value": "価格・価値", "cleanliness": "清潔さ",
    "ambience": "雰囲気・設備", "access": "アクセス", "booking": "予約・受付",
    "inclusivity": "食事制約・利用しやすさ", "operations": "案内・運営", "other": "その他",
}
SENTIMENTS = {"positive": "肯定", "negative": "否定", "mixed": "肯否混在", "neutral": "中立", "unavailable": "本文なし"}
CONFIDENCES = {"high": "高", "medium": "中", "low": "低"}
SKILL_DIR = Path(__file__).parent / "skills" / "google-review-report"


def write_json(path: Path, data):
    """Replace atomically; incomplete model output cannot masquerade as a saved batch."""
    path.parent.mkdir(parents=True, exist_ok=True)
    temp = path.with_suffix(path.suffix + ".tmp")
    temp.write_text(json.dumps(data, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    temp.replace(path)


def read_json(path: Path):
    return json.loads(path.read_text(encoding="utf-8"))


def source_data(path: Path) -> tuple[dict, list[dict], str]:
    raw = path.read_bytes()
    try:
        data = csv_source(raw) if path.suffix.lower() == ".csv" else json.loads(raw)
    except csv.Error as error:
        raise ValueError(f"CSVの形式を確認してください: {error}") from error
    if not isinstance(data, dict) or not isinstance(data.get("metadata"), dict) or not isinstance(data.get("reviews"), list):
        raise ValueError("収集CLIのJSON（metadata・reviews）が必要です。")
    reviews = data["reviews"]
    if not reviews:
        raise ValueError("口コミが0件です。収集を完了してから分析してください。")
    ids = []
    for index, review in enumerate(reviews):
        if not isinstance(review, dict) or not isinstance(review.get("review_id"), str) or not review["review_id"].strip():
            raise ValueError(f"口コミ{index + 1}: 口コミIDがありません。")
        ids.append(review["review_id"])
        for field in ("text", "owner_reply", "date_text", "raw_visible_text"):
            if not isinstance(review.get(field, ""), str):
                raise ValueError(f"口コミ{index + 1}: {field}は文字列で指定してください。")
        rating = review.get("rating")
        if rating is not None and (type(rating) is not int or rating not in range(1, 6)):
            raise ValueError(f"口コミ{index + 1}: 星評価が不正です。")
    if len(ids) != len(set(ids)):
        raise ValueError("元データに口コミIDの重複があります。重複を解消してから分析してください。")
    return data["metadata"], reviews, hashlib.sha256(raw).hexdigest()


def csv_source(raw: bytes) -> dict:
    """Import only the supplied CSV. No neighboring files or browser access."""
    from .cli import FIELDS
    try:
        text = raw.decode("utf-8-sig")
    except UnicodeDecodeError as error:
        raise ValueError("CSVをUTF-8形式で保存してください（BOM付きも利用できます）。") from error
    csv.field_size_limit(10_000_000)
    reader = csv.DictReader(io.StringIO(text, newline=""), strict=True)
    headers = reader.fieldnames
    if not headers or len(headers) != len(set(headers)):
        raise ValueError("CSVのヘッダーがないか、同名の列が重複しています。")
    aliases = {key: (title, key) for key, title in FIELDS}
    aliases["text"] += ("本文", "review_text")
    aliases["rating"] += ("評価", "stars")
    columns = {}
    for key, names in aliases.items():
        matches = [name for name in names if name in headers]
        if len(matches) > 1:
            raise ValueError(f"CSVに同じ項目の列が複数あります: {key}")
        if matches:
            columns[key] = matches[0]
    if not {"text", "rating"} <= set(columns):
        raise ValueError("CSVには「口コミ本文」「星評価」（またはtext・rating）の列が必要です。")
    reviews = []
    for index, values in enumerate(reader, 1):
        if None in values or any(value is None for value in values.values()):
            raise ValueError(f"CSVの{index}行目: 列数がヘッダーと一致しません。")
        if not any(values.values()):
            continue
        row = {key: values[column] for key, column in columns.items()}
        rating = row.get("rating", "").strip()
        if rating and not re.fullmatch(r"[1-5](?:\.0)?", rating):
            raise ValueError(f"CSVの{index}行目: 星評価は1〜5、または空欄です。")
        if not rating and not row["text"].strip():
            raise ValueError(f"CSVの{index}行目: 星評価または口コミ本文が必要です。")
        row["rating"] = int(float(rating)) if rating else None
        flag = row.get("text_may_be_truncated", "").strip().lower()
        if flag not in ("", "false", "0", "true", "1"):
            raise ValueError(f"CSVの{index}行目: 本文の省略ありはTrue/Falseで指定してください。")
        row["text_may_be_truncated"] = flag in ("true", "1")
        if not row.get("review_id", "").strip():
            row["review_id"] = f"csv-row-{index:06d}"
            row["review_id_origin"] = "CSV行番号（Googleの口コミIDは未提供）"
        reviews.append(row)
    places = {row.get("place_name", "").strip() for row in reviews} - {""}
    urls = {row.get("source_url", "").strip() for row in reviews} - {""}
    if len(places) > 1 or len(urls) > 1:
        raise ValueError("CSVに複数の店舗名または店舗URLがあります。施設ごとに分けてください。")
    collected = {row.get("collected_at", "").strip() for row in reviews} - {""}
    return {"metadata": {"input_format": "csv", "place_name": next(iter(places), "店舗名未提供"),
            "source_url": next(iter(urls), ""), "collected_at": max(collected) if collected else "未提供",
            "stop_reason": "利用者が指定したCSVを分析（Googleマップの全件取得はCSV単体では未確認）"}, "reviews": reviews}


def coverage(metadata: dict, reviews: list[dict]) -> dict:
    total = metadata.get("displayed_total_end")
    if type(total) is not int or total < 1:
        total = None
    count = len(reviews)
    verified = metadata.get("full_coverage_verified") is True and total == count
    return {"count": count, "displayed_total": total, "full_coverage_verified": verified,
            "input_format": metadata.get("input_format", "json"),
            "is_sample": metadata.get("sample") is True,
            "missing": max(0, total - count) if total else None,
            "text_count": sum(bool(r.get("text", "").strip()) for r in reviews),
            "truncated_count": sum(bool(r.get("text_may_be_truncated")) for r in reviews),
            "stop_reason": str(metadata.get("stop_reason", "不明"))}


def prepare(source: Path, output: Path, *, context: Path | None = None, batch_size: int = 20, reply_drafts: bool = False) -> Path:
    if not 1 <= batch_size <= 50:
        raise ValueError("分析の分割件数は1〜50件で指定してください。")
    metadata, reviews, digest = source_data(source)
    output = output.expanduser().resolve()
    output.mkdir(parents=True, exist_ok=True)
    if (output / "manifest.json").exists():
        raise ValueError("既存レポートを上書きしません。Skillに分析フォルダーを渡して続けるか、別の保存先を指定してください。")
    original = source.read_bytes()
    input_info = {}
    if source.suffix.lower() == ".csv":
        (output / "source.csv").write_bytes(original)
        normalized_source = (json.dumps({"metadata": metadata, "reviews": reviews}, ensure_ascii=False, indent=2) + "\n").encode("utf-8")
        digest = hashlib.sha256(normalized_source).hexdigest()
        (output / "source.json").write_bytes(normalized_source)
        input_info = {"input_format": "csv", "original_input_file": "source.csv", "original_input_sha256": hashlib.sha256(original).hexdigest()}
    else:
        (output / "source.json").write_bytes(original)
    if context:
        (output / "business-context.md").write_text(context.read_text(encoding="utf-8"), encoding="utf-8")
    text_reviews = []
    key_map = {}
    for i, review in enumerate(reviews, 1):
        key = f"R{i:04d}"
        key_map[key] = review["review_id"]
        if review.get("text", "").strip():
            # Text is untrusted evidence, never instructions. Poster names are unnecessary for analysis.
            text_reviews.append({"key": key, "rating": review.get("rating"), "date_text": review.get("date_text", ""),
                                 "text": review["text"], "owner_reply": review.get("owner_reply", ""),
                                 "text_may_be_truncated": bool(review.get("text_may_be_truncated")),
                                 "date_precision": review.get("date_precision", ""), "period_match": review.get("period_match", ""),
                                 "date_earliest": review.get("date_earliest", ""), "date_latest": review.get("date_latest", ""),
                                 "translation_visible": "Google による翻訳" in review.get("raw_visible_text", "") or "Googleによる翻訳" in review.get("raw_visible_text", "")})
    names = []
    for start in range(0, len(text_reviews), batch_size):
        name = f"batch-{len(names) + 1:03d}.json"
        write_json(output / "packets" / name, {"source_sha256": digest, "reviews": text_reviews[start:start + batch_size]})
        names.append(name)
    (output / "annotations").mkdir(exist_ok=True)
    manifest = {"schema_version": 1, "tool_version": __version__, "source_sha256": digest, "reply_drafts": reply_drafts,
                "place_name": str(metadata.get("place_name") or reviews[0].get("place_name") or "店舗"),
                "collected_at": str(metadata.get("collected_at", "不明")),
                "requested_url": str(metadata.get("requested_url", "")),
                "resolved_url": str(metadata.get("resolved_url", "")),
                "source_url": str(metadata.get("source_url", "")), "coverage": coverage(metadata, reviews),
                "keys": key_map, "batches": names, "prepared_at": datetime.now().astimezone().isoformat(timespec="seconds")}
    manifest.update(input_info)
    write_json(output / "manifest.json", manifest)
    return output


def workspace(output: Path):
    manifest = read_json(output / "manifest.json")
    if not isinstance(manifest, dict) or manifest.get("schema_version") != 1 or not isinstance(manifest.get("batches"), list) or any(not isinstance(name, str) or not re.fullmatch(r"batch-\d{3,}\.json", name) for name in manifest["batches"]) or len(manifest["batches"]) != len(set(manifest["batches"])):
        raise ValueError("manifest.jsonの形式または分割ファイル名が不正です。")
    metadata, reviews, digest = source_data(output / "source.json")
    if digest != manifest.get("source_sha256"):
        raise ValueError("source.jsonが準備後に変更されています。別の保存先でprepareをやり直してください。")
    if manifest.get("input_format") == "csv":
        if manifest.get("original_input_file") != "source.csv" or hashlib.sha256((output / "source.csv").read_bytes()).hexdigest() != manifest.get("original_input_sha256"):
            raise ValueError("入力CSVが準備後に変更されています。別の保存先でprepareをやり直してください。")
    keys = {f"R{i:04d}": r for i, r in enumerate(reviews, 1)}
    if manifest.get("keys") != {key: r["review_id"] for key, r in keys.items()}:
        raise ValueError("manifest.jsonの口コミ対応表が元データと一致しません。")
    manifest["coverage"] = coverage(metadata, reviews)
    return manifest, metadata, keys


def string(value, label: str, *, empty=False) -> str:
    if not isinstance(value, str) or (not empty and not value.strip()):
        raise ValueError(f"{label}: 文字列が必要です。")
    return value


def string_list(value, label: str) -> list[str]:
    if not isinstance(value, list) or any(not isinstance(item, str) or not item.strip() for item in value):
        raise ValueError(f"{label}: 文字列の配列が必要です。")
    return value


def normalized(value: str) -> str:
    return re.sub(r"\s+", " ", value).strip()


def validate_annotation(row: dict, raw: dict, key: str, *, reply_drafts=True):
    if not isinstance(row, dict) or row.get("key") != key:
        raise ValueError(f"{key}: keyが一致しません。")
    if row.get("sentiment") not in set(SENTIMENTS) - {"unavailable"}:
        raise ValueError(f"{key}: 本文の感情分類が不正です。")
    if row.get("confidence") not in CONFIDENCES:
        raise ValueError(f"{key}: confidenceはhigh/medium/lowです。")
    for field in ("summary_ja", "language") + (("reply_draft", "reply_language") if reply_drafts else ()):
        string(row.get(field), f"{key}.{field}")
    string_list(row.get("checks"), f"{key}.checks")
    aspects = row.get("aspects")
    if not isinstance(aspects, list):
        raise ValueError(f"{key}: aspectsが配列ではありません。")
    seen = set()
    for aspect in aspects:
        if not isinstance(aspect, dict) or aspect.get("category") not in CATEGORIES or aspect.get("polarity") not in set(SENTIMENTS) - {"unavailable"}:
            raise ValueError(f"{key}: カテゴリ・感情の分類が不正です。")
        if aspect["category"] in seen:
            raise ValueError(f"{key}: 同一カテゴリの重複があります。mixedに統合してください。")
        seen.add(aspect["category"])
        quote = string(aspect.get("quote"), f"{key}.quote")
        if normalized(quote) not in normalized(raw.get("text", "")):
            raise ValueError(f"{key}: 根拠引用が口コミ本文に存在しません: {quote[:70]}")
        string(aspect.get("detail"), f"{key}.detail")


def rating_only_annotation(key: str, raw: dict, *, reply_drafts=True) -> dict:
    rating = raw.get("rating")
    if rating and rating <= 2:
        reply = "評価をお寄せいただき、ありがとうございます。差し支えなければ、ご利用時に感じられた点をお聞かせください。"
    elif rating and rating >= 4:
        reply = "評価をお寄せいただき、ありがとうございます。またのご利用をお待ちしております。"
    else:
        reply = "評価をお寄せいただき、ありがとうございます。今後の参考として、ご感想もお聞かせいただけましたら幸いです。"
    return {"key": key, "sentiment": "unavailable", "confidence": "low", "language": "unknown",
            "summary_ja": "評価のみ。体験の内容・感情・原因は判断できません。", "aspects": [],
            "reply_draft": reply if reply_drafts else "", "reply_language": "ja" if reply_drafts else "",
            "checks": ["本文なし・返信言語と内容を確認", "定型の返信候補"] if reply_drafts else ["本文なし・体験の詳細は不明"]}


def load_annotations(output: Path, *, require_complete=True) -> tuple[dict, dict, list[dict]]:
    manifest, _metadata, keys = workspace(output)
    reply_drafts = manifest.get("reply_drafts", True)
    if type(reply_drafts) is not bool:
        raise ValueError("manifest.reply_draftsはtrue/falseで指定してください。")
    rows = {}
    for name in manifest["batches"]:
        path = output / "annotations" / name
        if not path.exists():
            if require_complete:
                raise ValueError(f"未分析の分割ファイルです: annotations/{name}")
            continue
        batch = read_json(path)
        if not isinstance(batch, dict) or batch.get("source_sha256") != manifest["source_sha256"] or not isinstance(batch.get("reviews"), list):
            raise ValueError(f"{name}: 元データのハッシュまたは形式が一致しません。")
        packet_keys = {r["key"] for r in read_json(output / "packets" / name)["reviews"]}
        batch_keys = []
        for row in batch["reviews"]:
            key = row.get("key") if isinstance(row, dict) else None
            if key not in packet_keys or key in rows:
                raise ValueError(f"{name}: 未知または重複したkeyです: {key}")
            validate_annotation(row, keys[key], key, reply_drafts=reply_drafts)
            rows[key] = row
            batch_keys.append(key)
        if set(batch_keys) != packet_keys:
            raise ValueError(f"{name}: 口コミが抜けています: {sorted(packet_keys - set(batch_keys))}")
    for key, raw in keys.items():
        if not raw.get("text", "").strip():
            rows[key] = rating_only_annotation(key, raw, reply_drafts=reply_drafts)
    if require_complete and set(rows) != set(keys):
        raise ValueError("本文の分析が全件揃っていません。")
    merged = []
    for key in keys:
        if key not in rows:
            continue
        row = dict(rows[key])
        if not reply_drafts:
            row.update(reply_draft="", reply_language="")
        row["raw"] = keys[key]
        row["checks"] = list(row["checks"])
        if keys[key].get("text_may_be_truncated"):
            row["checks"].append("原文の省略表示あり")
        if keys[key].get("date_precision"):
            row["checks"].append("日付: " + keys[key]["date_precision"] + " / " + keys[key].get("period_match", ""))
        if reply_drafts and keys[key].get("owner_reply"):
            row["checks"].append("既存返信あり・差し替え候補として確認")
        if row["confidence"] == "low" and keys[key].get("text", "").strip():
            row["checks"].append("分類の確信度が低い")
        merged.append(row)
    return manifest, keys, merged


def evidence(value, keys: dict, label: str):
    ids = string_list(value, label)
    if not ids or len(ids) != len(set(ids)) or any(key not in keys or not keys[key].get("text", "").strip() for key in ids):
        raise ValueError(f"{label}: 重複しない本文ありの口コミkeyが必要です。")


def load_synthesis(output: Path, manifest: dict, keys: dict) -> dict:
    data = read_json(output / "synthesis.json")
    if not isinstance(data, dict) or data.get("source_sha256") != manifest["source_sha256"]:
        raise ValueError("synthesis.jsonの元データのハッシュが一致しません。")
    for field in ("headline", "executive_summary", "business_type") + (("reply_policy",) if manifest.get("reply_drafts", True) else ()):
        string(data.get(field), f"synthesis.{field}")
    if not manifest.get("reply_drafts", True):
        data["reply_policy"] = ""
    for section in ("journey", "strengths", "concerns", "opportunities", "actions"):
        if not isinstance(data.get(section), list):
            raise ValueError(f"synthesis.{section}は配列で指定してください。")
        action_ids = set()
        for i, item in enumerate(data[section]):
            label = f"{section}[{i}]"
            if not isinstance(item, dict):
                raise ValueError(f"{label}: オブジェクトが必要です。")
            evidence(item.get("evidence_ids"), keys, label)
            fields = ("id", "title", "detail", "owner", "timeframe", "priority", "priority_reason", "effort", "kpi", "baseline", "target") if section == "actions" else ("title", "observation")
            if section == "journey":
                fields = ("stage", "observation", "friction", "opportunity")
            if section == "concerns":
                fields += ("hypothesis", "alternative", "verify")
            for field in fields:
                string(item.get(field), f"{label}.{field}")
            if section == "actions":
                if item["id"] in action_ids or item["priority"] not in ("高", "中", "低") or item["timeframe"] not in ("0〜7日", "8〜30日", "31〜90日"):
                    raise ValueError(f"{label}: アクションID・優先度・期間が不正です。")
                action_ids.add(item["id"])
    return data


def recency(date_text: str) -> str:
    # Displayed dates are approximate and edits do not establish original posting dates.
    if any(word in date_text for word in ("編集", "edited", "Edited")):
        return "編集日表示・別集計"
    units = {"秒": 1/86400, "分": 1/1440, "時間": 1/24, "日": 1, "週": 7,
             "second": 1/86400, "minute": 1/1440, "hour": 1/24, "day": 1, "week": 7}
    recent = re.search(r"(\d+)\s*(秒|分|時間|日|週)(?:間)?前", date_text) or re.search(r"(\d+)\s*(second|minute|hour|day|week)s? ago", date_text, re.I)
    if recent:
        days = int(recent[1]) * units[recent[2].lower()]
        return "3か月以内と表示" if days <= 93 else "4〜11か月と表示" if days < 365 else "1年以上と表示"
    match = re.search(r"(\d+)\s*(?:か月|ヶ月|ケ月|月)前", date_text) or re.search(r"(\d+) months? ago", date_text, re.I)
    if match:
        months = int(match[1])
        return "3か月以内と表示" if months <= 3 else "4〜11か月と表示" if months < 12 else "1年以上と表示"
    if re.search(r"\d+\s*年前|\d+ years? ago", date_text, re.I):
        return "1年以上と表示"
    return "時期を判定できない"


def statistics(manifest: dict, rows: list[dict]) -> dict:
    rating_values = [r["raw"].get("rating") for r in rows if r["raw"].get("rating") in range(1, 6)]
    text = [r for r in rows if r["sentiment"] != "unavailable"]
    categories = []
    for code, title in CATEGORIES.items():
        counts = Counter(a["polarity"] for r in text for a in r["aspects"] if a["category"] == code)
        total = sum(counts.values())
        categories.append({"code": code, "title": title, "mentions": total,
                           "share_text_pct": round(100 * total / len(text), 1) if text else 0,
                           **{s: counts[s] for s in ("positive", "negative", "mixed", "neutral")},
                           "negative_or_mixed": counts["negative"] + counts["mixed"]})
    cohorts = []
    for name in ("3か月以内と表示", "4〜11か月と表示", "1年以上と表示", "編集日表示・別集計", "時期を判定できない"):
        group = [r for r in rows if recency(r["raw"].get("date_text", "")) == name]
        values = [r["raw"].get("rating") for r in group if r["raw"].get("rating") in range(1, 6)]
        cohorts.append({"label": name, "count": len(group), "text_count": sum(r["sentiment"] != "unavailable" for r in group),
                        "rating_count": len(values), "mean": round(sum(values) / len(values), 2) if values else None})
    return {"count": len(rows), "text_count": len(text), "rating_only_count": len(rows) - len(text),
            "rating_count": len(rating_values), "mean": round(sum(rating_values) / len(rating_values), 2) if rating_values else None,
            "stars": {str(i): rating_values.count(i) for i in range(1, 6)},
            "sentiments": {s: sum(r["sentiment"] == s for r in rows) for s in SENTIMENTS},
            "categories": categories, "cohorts": cohorts,
            "cross": {str(i): {s: sum(r["raw"].get("rating") == i and r["sentiment"] == s for r in rows) for s in SENTIMENTS} for i in range(1, 6)},
            "high_star_concerns": sum((r["raw"].get("rating") or 0) >= 4 and any(a["polarity"] in ("negative", "mixed") for a in r["aspects"]) for r in rows),
            "coverage": manifest["coverage"]}


def write_sheet(book: Workbook, name: str, headers: list[str], rows):
    sheet = book.create_sheet(name)
    sheet.append(headers)
    for row in rows:
        sheet.append(row)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    sheet.sheet_view.zoomScale = 85
    sheet.row_dimensions[1].height = 36
    for column, header in enumerate(headers, 1):
        width = 60 if any(word in header for word in ("本文", "引用", "解釈", "内容", "要約", "理由", "仮説", "確認方法", "機会", "返信案")) else 28 if any(word in header for word in ("URL", "日時", "日付", "カテゴリ", "口コミID")) else 18
        sheet.column_dimensions[get_column_letter(column)].width = width
    for row in sheet:
        if row[0].row > 1:
            lines = max((len(str(cell.value or "")) * 2 // sheet.column_dimensions[cell.column_letter].width + 1 for cell in row), default=1)
            sheet.row_dimensions[row[0].row].height = min(150, max(30, lines * 15))
        for cell in row:
            # Original review/model strings are literal Excel text, including
            # leading '='. Preserve text without allowing formula execution.
            if isinstance(cell.value, str):
                cell.data_type = "s"
            cell.font = Font(name="Arial", size=10, color="253444")
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            if cell.row == 1:
                cell.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="263F52")
            elif cell.row % 2 == 0:
                cell.fill = PatternFill("solid", fgColor="F2F5F7")
            if cell.row > 1 and isinstance(cell.value, float):
                cell.number_format = "0.00"
    if name == "改善アクション" and sheet.max_row > 1:
        column = get_column_letter(sheet.max_column)
        states = DataValidation(type="list", formula1='"未着手・事業者確認,進行中,完了,保留"')
        sheet.add_data_validation(states)
        states.add(f"{column}2:{column}{sheet.max_row}")
        for row in range(2, sheet.max_row + 1):
            sheet[f"{column}{row}"].fill = PatternFill("solid", fgColor="FFF1CC")
    return sheet


def render(output: Path) -> Path:
    manifest, keys, rows = load_annotations(output)
    synthesis = load_synthesis(output, manifest, keys)
    stats = statistics(manifest, rows)
    generated_at = datetime.now().astimezone().isoformat(timespec="seconds")
    write_json(output / "statistics.json", stats)
    write_json(output / "analysis.json", {"source_sha256": manifest["source_sha256"], "generated_at": generated_at,
                                         "reviews": [{k: v for k, v in r.items() if k != "raw"} for r in rows], "synthesis": synthesis})
    book = Workbook()
    book.remove(book.active)
    write_sheet(book, "概要", ["項目", "内容"], [
        ["施設名", manifest["place_name"]], ["取得元URL", manifest["source_url"]],
        ["入力URL", manifest.get("requested_url", "")], ["取得日時", manifest["collected_at"]],
        ["生成日時", generated_at], ["保存件数", stats["count"]], ["画面の総件数", manifest["coverage"]["displayed_total"]],
        ["本文あり", stats["text_count"]], ["評価のみ", stats["rating_only_count"]], ["平均星", stats["mean"]],
        ["分析対象", "指定CSVの全行" if manifest["coverage"].get("input_format") == "csv" else "指定された元データの全行"],
        ["Google全件照合", "照合済み" if manifest["coverage"]["full_coverage_verified"] else "未確認（CSV単体では確認不可）" if manifest["coverage"].get("input_format") == "csv" else "未確認・取得した範囲の分析"],
        ["停止理由", manifest["coverage"]["stop_reason"]], ["返信案", "生成あり" if manifest.get("reply_drafts", True) else "生成なし"],
        ["全体概要", synthesis["executive_summary"]], ["業種", synthesis["business_type"]],
        ["利用方法", "管理IDで原文・根拠・改善案を照合できます。改善アクションの状態列は実行状況に合わせて変更してください。集計は生成時点の値です。"],
        ["原文SHA256", manifest["source_sha256"]],
        ["入力CSVのSHA256", manifest.get("original_input_sha256", "未提供")],
    ] + [[f"星{i}の件数", stats["stars"][str(i)]] for i in range(5, 0, -1)]
      + [[f"本文感情：{title}", stats["sentiments"][code]] for code, title in SENTIMENTS.items()]
      + [[f"{group['label']}：件数", group["count"]] for group in stats["cohorts"]])
    write_sheet(book, "分類済み口コミ", ["管理ID", "口コミID", "投稿者", "星", "表示日付", "本文", "本文感情", "カテゴリ", "要約", "確信度", "要確認", "既存返信", "収集日時", "日付範囲の開始", "日付範囲の終了", "日付の精度", "期間判定"],
               ([r["key"], r["raw"]["review_id"], r["raw"].get("author", ""), r["raw"].get("rating"), r["raw"].get("date_text", ""), r["raw"].get("text", ""), SENTIMENTS[r["sentiment"]], " / ".join(CATEGORIES[a["category"]] for a in r["aspects"]), r["summary_ja"], CONFIDENCES[r["confidence"]], " / ".join(r["checks"]), r["raw"].get("owner_reply", ""), manifest["collected_at"]] + [r["raw"].get(key, "") for key in ("date_earliest", "date_latest", "date_precision", "period_match")] for r in rows))
    write_sheet(book, "話題別根拠", ["管理ID", "口コミID", "カテゴリ", "感情", "根拠引用", "解釈"],
               ([r["key"], r["raw"]["review_id"], CATEGORIES[a["category"]], SENTIMENTS[a["polarity"]], a["quote"], a["detail"]] for r in rows for a in r["aspects"]))
    if manifest.get("reply_drafts", True):
        write_sheet(book, "返信案", ["管理ID", "口コミID", "投稿者", "星", "本文", "要確認", "既存返信", "返信言語", "返信案", "状態"],
               ([r["key"], r["raw"]["review_id"], r["raw"].get("author", ""), r["raw"].get("rating"), r["raw"].get("text", ""), " / ".join(r["checks"]), r["raw"].get("owner_reply", ""), r["reply_language"], r["reply_draft"], "既存返信あり・変更は要確認" if r["raw"].get("owner_reply") else "下書き・投稿前確認"] for r in rows))
    write_sheet(book, "カテゴリ集計", ["カテゴリ", "言及口コミ数", "本文あり口コミに占める割合%", "肯定", "否定", "肯否混在", "中立", "否定または混在", "本文あり母数"],
               ([c["title"], c["mentions"], c["share_text_pct"], c["positive"], c["negative"], c["mixed"], c["neutral"], c["negative_or_mixed"], len([r for r in rows if r["sentiment"] != "unavailable"])] for c in stats["categories"]))
    write_sheet(book, "改善アクション", ["アクションID", "改善案", "実行内容", "優先度", "優先度の理由", "担当案", "着手からの期間案", "工数案", "検証指標", "現状値", "目標案", "根拠口コミ", "状態"],
               ([a["id"], a["title"], a["detail"], a["priority"], a["priority_reason"], a["owner"], a["timeframe"], a["effort"], a["kpi"], a["baseline"], a["target"], " / ".join(a["evidence_ids"]), "未着手・事業者確認"] for a in synthesis["actions"]))
    write_sheet(book, "分析と提案", ["区分", "項目", "観察・内容", "原因の仮説", "別の説明", "確認方法", "摩擦・確認点", "機会", "根拠口コミ"],
                ([label, item.get("title", item.get("stage", "")), item.get("observation", ""), item.get("hypothesis", ""),
                  item.get("alternative", ""), item.get("verify", ""), item.get("friction", ""), item.get("opportunity", ""), " / ".join(item["evidence_ids"])]
                 for section, label in [("journey", "顧客体験"), ("strengths", "強み"), ("concerns", "課題"), ("opportunities", "事業機会")]
                 for item in synthesis[section]))
    temporary = output / "口コミ分析.tmp.xlsx"
    book.save(temporary)
    book.close()
    temporary.replace(output / "口コミ分析.xlsx")
    safe_rows = []
    for r in rows:
        safe_rows.append({k: v for k, v in r.items() if k != "raw"} | {"rating": r["raw"].get("rating"), "date": r["raw"].get("date_text", ""), "text": r["raw"].get("text", ""), "owner_reply": r["raw"].get("owner_reply", "")})
    payload = {"manifest": {k: v for k, v in manifest.items() if k != "keys"}, "stats": stats, "synthesis": synthesis,
               "rows": safe_rows, "categories": CATEGORIES, "sentiments": SENTIMENTS, "generated_at": generated_at}
    template = (Path(__file__).parent / "report_assets" / "report.html").read_text(encoding="utf-8")
    script = (Path(__file__).parent / "report_assets" / "report.js").read_text(encoding="utf-8")
    css = (Path(__file__).parent / "report_assets" / "report.css").read_text(encoding="utf-8")
    safe_json = json.dumps(payload, ensure_ascii=False).replace("&", "\\u0026").replace("<", "\\u003c").replace(">", "\\u003e").replace("\u2028", "\\u2028").replace("\u2029", "\\u2029")
    replacements = {"TITLE": html.escape(manifest["place_name"] + " 口コミ分析レポート"), "CSS": css, "DATA": safe_json, "JS": script}
    document = re.sub(r"\{\{(TITLE|CSS|DATA|JS)\}\}", lambda match: replacements[match[1]], template)
    path = output / "口コミレポート.html"
    path.write_text(document, encoding="utf-8")
    return path
