import json
import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook, load_workbook
from google_maps_reviews.cli import FIELDS, export_reviews, write_review_excel
from google_maps_reviews.excel import display_lines, split_display_text, write_table


class ExcelLayoutTest(unittest.TestCase):
    def fixture(self):
        def review(key, text):
            return {"review_id": key, "place_name": "架空店舗", "author": "架空投稿者", "rating": 4,
                    "date_text": "1年前", "text": text, "owner_reply": "既存の返信\n原文を保持",
                    "source_url": "https://www.google.com/maps/test", "period_match": "テスト"}
        return {"metadata": {"place_name": "架空店舗", "count": 2, "scanned_count": 5,
                "scanned_text_review_count": 4, "excluded_rating_only_count": 1,
                "date_from": "2023-01-01", "date_to": "2026-09-30", "period_uncertain_count": 1,
                "period_excluded_count": 1, "scan_coverage_verified": True, "incomplete": True,
                "review_filter": "text_only"},
                "reviews": [review("a", "短い本文"), review("b", "料理と接客の口コミです。\n" * 12)],
                "uncertain_reviews": [review("c", "日付の確認が必要。\n" * 12)],
                "excluded_reviews": [review("d", "以前の口コミ本文。\n" * 12)]}

    def test_all_review_groups_preserve_values_and_receive_the_same_layout(self):
        data = self.fixture()
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "reviews.xlsx"
            write_review_excel(data, path)
            book = load_workbook(path)
            self.assertEqual(book.active.title, "取得情報")
            for key, name in (("reviews", "口コミ"), ("uncertain_reviews", "期間境界・日付不明"), ("excluded_reviews", "期間外")):
                sheet = book[name]
                self.assertEqual([cell.value for cell in sheet[1]], [title for _, title in FIELDS])
                self.assertEqual(sheet.freeze_panes, "E2")
                self.assertEqual(sheet.auto_filter.ref, f"A1:Q{len(data[key])+1}")
                self.assertTrue(sheet.column_dimensions["G"].hidden)
                self.assertTrue(sheet.column_dimensions["M"].hidden)
                self.assertTrue(sheet.column_dimensions["N"].hidden)
                self.assertFalse(sheet.column_dimensions["E"].hidden)
                for index, original in enumerate(data[key], 2):
                    self.assertEqual(sheet.cell(index, 5).value, original["text"])
                    self.assertEqual(sheet.cell(index, 6).value, original["owner_reply"])
                    self.assertEqual(sheet.cell(index, 7).value, original["review_id"])
                    self.assertTrue(sheet.cell(index, 5).alignment.wrap_text)
                    self.assertEqual(sheet.cell(index, 5).font.name, "Arial")
                    self.assertLessEqual(sheet.row_dimensions[index].height, 409)
                self.assertEqual(sheet.page_setup.orientation, "landscape")
                self.assertEqual(sheet.page_setup.fitToWidth, 1)
                self.assertEqual(sheet.print_title_rows, "$1:$1")
            self.assertLess(book["口コミ"].row_dimensions[2].height, book["口コミ"].row_dimensions[3].height)
            self.assertGreater(book["期間境界・日付不明"].row_dimensions[2].height, 150)
            info = {row[0]: row[1] for row in book["取得情報"].iter_rows(min_row=2, values_only=True)}
            self.assertEqual(info["保存件数"], 2)
            self.assertIsInstance(info["保存件数"], int)
            self.assertEqual(info["保存対象"], "本文ありの口コミのみ")
            self.assertEqual(info["一覧全件の読取件数照合"], "確認済み")
            book.close()

    def test_newlines_and_unicode_are_preserved_across_long_text_chunks(self):
        body = "日本語 😀 e\u0301\tコメント\r\n\n" * 4000
        for text in (body, "\n" * 85, "日本語" * 5000):
            chunks = split_display_text(text, 88)
            self.assertEqual("".join(chunks), text)
            self.assertTrue(all(display_lines(chunk, 88) <= 21 for chunk in chunks))
        with tempfile.TemporaryDirectory() as folder:
            data = self.fixture()
            data["uncertain_reviews"][0]["text"] = body
            path = Path(folder) / "long.xlsx"
            write_review_excel(data, path)
            book = load_workbook(path)
            self.assertEqual(book.sheetnames[-1], "長文の続き")
            chunks = [row[4] for row in book["長文の続き"].iter_rows(min_row=2, values_only=True)
                      if row[0] == "期間境界・日付不明" and row[1] == "E2"]
            self.assertEqual("".join(chunks), body)
            self.assertIsNone(book["期間境界・日付不明"]["E2"].hyperlink)
            self.assertIsNotNone(book["期間境界・日付不明"]["R2"].hyperlink)
            self.assertTrue(all(row.height <= 409 for row in book["長文の続き"].row_dimensions.values()))
            book.close()

    def test_hidden_technical_text_does_not_make_short_reviews_tall(self):
        book = Workbook()
        book.remove(book.active)
        sheet = write_table(book, "テスト", ["本文", "原表示"], [["短い本文", "\n" * 500]],
                            widths=[72, 30], hidden_columns=(2,))
        self.assertEqual(sheet.row_dimensions[2].height, 30)
        self.assertNotIn("長文の続き", book.sheetnames)
        self.assertEqual(sheet["B2"].value, "\n" * 500)

    def test_layout_export_does_not_change_csv_or_json_review_evidence(self):
        body = '=HYPERLINK("https://example.com")\n日本語\n\n末尾'
        review = dict(self.fixture()["reviews"][0], text=body)
        with tempfile.TemporaryDirectory() as folder:
            paths = export_reviews([review], {"count": 1, "review_filter": "text_only"}, Path(folder), "test")
            original = json.loads(paths[2].read_text())
            self.assertEqual(original["reviews"], [review])
            self.assertIn(body.encode('utf-8').replace(b'"', b'""'), paths[0].read_bytes())
            book = load_workbook(paths[1])
            self.assertEqual(book["口コミ"]["E2"].value, body)
            self.assertFalse(any(cell.data_type == "f" for sheet in book for row in sheet for cell in row))
            book.close()


if __name__ == "__main__":
    unittest.main()
