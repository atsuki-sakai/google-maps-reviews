"""Analysis workbooks keep keyed source data readable and recoverable."""
import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook, load_workbook

from google_maps_reviews.reporting import write_sheet


REVIEW_HEADERS = [
    "管理ID", "口コミID", "投稿者", "星", "表示日付", "本文", "本文感情", "カテゴリ",
    "要約", "確信度", "要確認", "既存返信", "収集日時", "日付範囲の開始",
    "日付範囲の終了", "日付の精度", "期間判定",
]
ACTION_HEADERS = [
    "アクションID", "改善案", "実行内容", "優先度", "優先度の理由", "担当案",
    "着手からの期間案", "工数案", "検証指標", "現状値", "目標案", "根拠口コミ", "状態",
]


class ReportingLayoutTest(unittest.TestCase):
    def new_book(self):
        book = Workbook()
        book.remove(book.active)
        self.addCleanup(book.close)
        return book

    def test_classification_keeps_primary_keys_visible_and_original_columns_available(self):
        book = self.new_book()
        values = ["R0001", "review-1", "=AUTHOR()", 5, "2か月前", "原文の本文", "肯定",
                  "接客", "要約", "高", "", "既存返信", "2026-10-06T12:00:00+07:00",
                  "2026-07-01", "2026-08-31", "estimated", "within"]
        sheet = write_sheet(book, "分類済み口コミ", REVIEW_HEADERS, [values])
        self.assertEqual([cell.value for cell in sheet[1]], REVIEW_HEADERS)
        self.assertEqual([cell.value for cell in sheet[2]], values)
        self.assertEqual(sheet.freeze_panes, "D2")
        self.assertFalse(sheet.column_dimensions["A"].hidden)
        self.assertFalse(sheet.column_dimensions["F"].hidden)
        self.assertTrue(sheet.column_dimensions["B"].hidden)
        self.assertTrue(sheet.column_dimensions["M"].hidden)
        self.assertFalse(sheet.column_dimensions["Q"].hidden)
        self.assertGreater(sheet.column_dimensions["F"].width, sheet.column_dimensions["D"].width)
        self.assertEqual(sheet.auto_filter.ref, "A1:Q2")
        self.assertEqual(sheet["C2"].data_type, "s")
        self.assertEqual(sheet["D2"].data_type, "n")

    def test_explicit_line_breaks_receive_enough_height_in_short_report_rows(self):
        book = self.new_book()
        explanation = "\n".join(f"・改善の根拠 {number}" for number in range(12))
        sheet = write_sheet(book, "概要", ["項目", "内容"], [["全体概要", explanation]])
        self.assertEqual(sheet["B2"].value, explanation)
        self.assertGreater(sheet.row_dimensions[2].height, 150)
        self.assertLessEqual(sheet.row_dimensions[2].height, 409)
        self.assertTrue(sheet["B2"].alignment.wrap_text)
        self.assertEqual(sheet["B2"].alignment.vertical, "top")

    def test_long_original_and_reply_are_recoverable_without_moving_review_keys(self):
        book = self.new_book()
        original = "\n\n".join(f"・{number}：提供時間と料理の感想。\n次の行の説明。" for number in range(90))
        reply = "=これは既存返信の原文です。\n" + "確認事項を説明する行。\n" * 70
        values = ["R0001", "review-1", "投稿者", 4, "1年前", original, "肯否混在", "品質 / 接客",
                  "原文に基づく要約", "高", "既存返信あり", reply, "", "", "", "", ""]
        write_sheet(book, "分類済み口コミ", REVIEW_HEADERS, [values])
        with tempfile.TemporaryDirectory() as root:
            path = Path(root) / "report.xlsx"
            book.save(path)
            saved = load_workbook(path)
            try:
                sheet = saved["分類済み口コミ"]
                self.assertEqual((sheet["A2"].value, sheet["B2"].value), ("R0001", "review-1"))
                self.assertEqual(sheet["F2"].value, original)
                self.assertEqual(sheet["L2"].value, reply)
                self.assertEqual(sheet.max_row, 2)
                continuation = saved["長文の続き"]
                for coordinate, expected in (("F2", original), ("L2", reply)):
                    chunks = sorted((row[3], row[4]) for row in continuation.iter_rows(min_row=2, values_only=True)
                                    if row[0] == "分類済み口コミ" and row[1] == coordinate)
                    self.assertTrue(chunks)
                    self.assertEqual("".join(text for _, text in chunks), expected)
                self.assertFalse(sheet["F2"].hyperlink)
                self.assertTrue(sheet["R2"].hyperlink)
                self.assertFalse(sheet["R2"].hyperlink.target)
                for worksheet in saved:
                    self.assertTrue(all((dimension.height or 0) <= 409 for dimension in worksheet.row_dimensions.values()))
                    self.assertFalse(any(cell.data_type == "f" for row in worksheet for cell in row))
            finally:
                saved.close()

    def test_action_status_input_and_evidence_keys_survive_multiline_explanations(self):
        book = self.new_book()
        detail = "\n".join(f"・実行手順 {number} の確認" for number in range(12))
        values = ["A01", "提供時間の記録", detail, "高", "本文の指摘", "店長", "7日", "低",
                  "提供時間", "未測定", "まず計測", "R0001 / R0002", "未着手・事業者確認"]
        sheet = write_sheet(book, "改善アクション", ACTION_HEADERS, [values, ["A02", *values[1:]]])
        self.assertEqual(sheet.freeze_panes, "C2")
        self.assertEqual(sheet["C2"].value, detail)
        self.assertEqual(sheet["L2"].value, "R0001 / R0002")
        self.assertEqual(sheet.max_row, 3)
        validations = list(sheet.data_validations.dataValidation)
        self.assertEqual(len(validations), 1)
        self.assertEqual(str(validations[0].sqref), "M2:M3")
        self.assertIn("進行中", validations[0].formula1)
        self.assertGreater(sheet.row_dimensions[2].height, 150)

    def test_long_action_keeps_validation_on_status_after_navigation_column(self):
        book = self.new_book()
        detail = "実行手順の説明。\n" * 60
        values = ["A01", "改善案", detail, "高", "根拠", "店長", "7日", "低",
                  "指標", "未測定", "まず計測", "R0001", "未着手・事業者確認"]
        sheet = write_sheet(book, "改善アクション", ACTION_HEADERS, [values])
        self.assertEqual(sheet["N1"].value, "全文の参照")
        self.assertEqual(str(sheet.data_validations.dataValidation[0].sqref), "M2")
        self.assertTrue(sheet["N2"].hyperlink)
        self.assertFalse(sheet["C2"].hyperlink)


if __name__ == "__main__":
    unittest.main()
