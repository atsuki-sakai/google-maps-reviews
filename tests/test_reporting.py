import io
import csv
import json
import re
import sys
import tempfile
import unittest
from contextlib import redirect_stderr, redirect_stdout
from pathlib import Path
from unittest.mock import patch

from openpyxl import load_workbook
from google_maps_reviews import reporting as r
from google_maps_reviews import report_cli, cli


class ReportingTest(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.source = self.root / "reviews.json"
        self.output = self.root / "report"
        self.raw = [
            {"review_id": "review-1", "author": "=AUTHOR()", "rating": 5, "date_text": "2 か月前", "text": "料理は美味しい。ただ、提供は遅かった。", "owner_reply": "既存の返信"},
            {"review_id": "review-2", "rating": 1, "date_text": "最終編集: 1 年前", "text": "清潔で良かった。", "owner_reply": ""},
            {"review_id": "review-3", "rating": 4, "date_text": "1 年前", "text": "", "owner_reply": ""},
        ]
        self.save_source()
        r.prepare(self.source, self.output, reply_drafts=True)
        self.sha = r.read_json(self.output / "manifest.json")["source_sha256"]
        self.annotations = [
            {"key": "R0001", "sentiment": "mixed", "confidence": "high", "language": "ja", "summary_ja": "味は好評、待ち時間は不満", "aspects": [
                {"category": "quality", "polarity": "positive", "quote": "料理は美味しい", "detail": "料理の味を肯定"},
                {"category": "waiting", "polarity": "negative", "quote": "提供は遅かった", "detail": "提供速度を否定"}],
             "reply_draft": "ご感想をありがとうございます。提供までお待たせし申し訳ありません。", "reply_language": "ja", "checks": []},
            {"key": "R0002", "sentiment": "positive", "confidence": "medium", "language": "ja", "summary_ja": "清潔さを評価", "aspects": [
                {"category": "cleanliness", "polarity": "positive", "quote": "清潔で良かった", "detail": "清潔さを肯定"}],
             "reply_draft": "清潔さへのご感想をありがとうございます。", "reply_language": "ja", "checks": ["星と本文の評価が異なる"]},
        ]
        self.save_annotations()
        self.synthesis = {"source_sha256": self.sha, "headline": "品質を維持し、提供の確認を", "executive_summary": "料理を評価する口コミに提供の遅さの記述がある。", "business_type": "飲食店", "reply_policy": "未確認の改善を約束しない。", "journey": [], "strengths": [], "concerns": [], "opportunities": [], "actions": []}
        r.write_json(self.output / "synthesis.json", self.synthesis)

    def save_source(self, **metadata):
        r.write_json(self.source, {"metadata": {"place_name": "テスト店舗", "displayed_total_end": 3, "full_coverage_verified": True, **metadata}, "reviews": self.raw})

    def save_annotations(self):
        r.write_json(self.output / "annotations" / "batch-001.json", {"source_sha256": self.sha, "reviews": self.annotations})

    def test_star_sentiment_and_rating_only_denominators_are_separate(self):
        manifest, _, rows = r.load_annotations(self.output)
        stats = r.statistics(manifest, rows)
        self.assertEqual((stats["count"], stats["text_count"], stats["rating_only_count"]), (3, 2, 1))
        self.assertEqual(stats["mean"], 3.33)
        self.assertEqual(stats["sentiments"]["positive"], 1)  # one-star review has positive text
        self.assertEqual(stats["cross"]["1"]["positive"], 1)
        self.assertEqual(stats["high_star_concerns"], 1)
        self.assertEqual(stats["categories"][0]["share_text_pct"], 50)
        self.assertEqual(rows[-1]["sentiment"], "unavailable")
        self.assertEqual(rows[-1]["aspects"], [])
        self.assertIn("既存返信", " ".join(rows[0]["checks"]))
        self.assertEqual(sum(c["count"] for c in stats["cohorts"]), 3)

    def test_fabricated_quote_blocks_generation(self):
        self.annotations[0]["aspects"][0]["quote"] = "笑顔で接客してくれた"
        self.save_annotations()
        with self.assertRaisesRegex(ValueError, "R0001.*本文に存在"):
            r.render(self.output)
        self.assertFalse((self.output / "口コミレポート.html").exists())

    def test_missing_unknown_and_duplicate_rows_block_generation(self):
        for rows in (self.annotations[:1], self.annotations + [self.annotations[0]], [{**self.annotations[0], "key": "R9999"}, self.annotations[1]]):
            r.write_json(self.output / "annotations" / "batch-001.json", {"source_sha256": self.sha, "reviews": rows})
            with self.assertRaises(ValueError):
                r.load_annotations(self.output)

    def test_duplicate_category_is_not_double_counted(self):
        self.annotations[0]["aspects"].append(self.annotations[0]["aspects"][0])
        self.save_annotations()
        with self.assertRaisesRegex(ValueError, "同一カテゴリ"):
            r.load_annotations(self.output)

    def test_wrong_dataset_hash_and_tampered_source_are_rejected(self):
        r.write_json(self.output / "annotations" / "batch-001.json", {"source_sha256": "other", "reviews": self.annotations})
        with self.assertRaises(ValueError):
            r.load_annotations(self.output)
        self.save_annotations()
        data = r.read_json(self.output / "source.json")
        data["reviews"][0]["text"] = "変更後"
        r.write_json(self.output / "source.json", data)
        with self.assertRaisesRegex(ValueError, "変更"):
            r.load_annotations(self.output)

    def test_synthesis_cannot_use_rating_only_as_cause_evidence(self):
        self.synthesis["strengths"] = [{"title": "推論", "observation": "原因", "evidence_ids": ["R0003"]}]
        r.write_json(self.output / "synthesis.json", self.synthesis)
        with self.assertRaisesRegex(ValueError, "本文あり"):
            r.render(self.output)

    def test_partial_source_remains_partial_and_truncated_flag_is_visible(self):
        self.raw[0]["text_may_be_truncated"] = True
        self.save_source(displayed_total_end=10)
        other = self.root / "partial"
        r.prepare(self.source, other)
        manifest = r.read_json(other / "manifest.json")
        self.assertFalse(manifest["coverage"]["full_coverage_verified"])
        self.assertEqual(manifest["coverage"]["missing"], 7)
        self.assertEqual(manifest["coverage"]["truncated_count"], 1)

    def test_relative_dates_do_not_turn_edits_into_posting_dates(self):
        self.assertEqual(r.recency("最終編集: 2 日前"), "編集日表示・別集計")
        self.assertEqual(r.recency("2 か月前"), "3か月以内と表示")
        self.assertEqual(r.recency("5 months ago"), "4〜11か月と表示")
        self.assertEqual(r.recency("2025/10/01"), "時期を判定できない")
        self.assertEqual(r.recency("20 weeks ago"), "4〜11か月と表示")

    def test_display_coverage_is_recomputed_from_original_metadata(self):
        manifest = r.read_json(self.output / "manifest.json")
        manifest["coverage"]["count"] = 999
        r.write_json(self.output / "manifest.json", manifest)
        self.assertEqual(r.workspace(self.output)[0]["coverage"]["count"], 3)

    def test_manifest_cannot_reference_files_outside_packet_directory(self):
        manifest = r.read_json(self.output / "manifest.json")
        manifest["batches"] = ["../../other.json"]
        r.write_json(self.output / "manifest.json", manifest)
        with self.assertRaisesRegex(ValueError, "分割ファイル名"):
            r.workspace(self.output)

    def test_offline_report_neutralizes_html_and_preserves_template_like_data(self):
        malicious = '</script><script>window.pwned=1</script>{{JS}}'
        self.annotations[0]["reply_draft"] = malicious
        self.annotations[0]["summary_ja"] = "=FORMULA()"
        self.save_annotations()
        document = r.render(self.output).read_text()
        self.assertNotIn("<script>window.pwned", document)
        payload = json.loads(re.search(r'<script id="report-data" type="application/json">(.*?)</script>', document, re.S)[1])
        self.assertEqual(payload["rows"][0]["reply_draft"], malicious)
        self.assertNotIn("author", payload["rows"][0])
        book = load_workbook(self.output / "口コミ分析.xlsx")
        sheet = book["分類済み口コミ"]
        self.assertEqual(sheet["F2"].value, self.raw[0]["text"])
        self.assertEqual(sheet["C2"].value, "=AUTHOR()")
        self.assertEqual(sheet["I2"].value, "=FORMULA()")
        self.assertEqual(sheet["C2"].data_type, "s")
        self.assertEqual(sheet["I2"].data_type, "s")
        self.assertEqual(sheet.freeze_panes, "A2")
        self.assertEqual(sheet.auto_filter.ref, "A1:Q4")
        self.assertEqual(book["返信案"].max_row, 4)
        self.assertEqual(book["カテゴリ集計"]["B2"].data_type, "n")
        book.close()
        self.assertEqual(len(list(self.output.glob("*.csv"))), 0)
        self.assertEqual(len(list(self.output.glob("*.xlsx"))), 1)

    def test_analysis_without_replies_needs_no_reply_fields_or_reply_sheet(self):
        manifest = r.read_json(self.output / "manifest.json")
        manifest["reply_drafts"] = False
        r.write_json(self.output / "manifest.json", manifest)
        for row in self.annotations:
            row.pop("reply_draft")
            row.pop("reply_language")
        self.save_annotations()
        self.synthesis.pop("reply_policy")
        r.write_json(self.output / "synthesis.json", self.synthesis)
        r.render(self.output)
        _, _, rows = r.load_annotations(self.output)
        self.assertTrue(all(row["reply_draft"] == "" for row in rows))
        self.assertNotIn("返信", " ".join(rows[-1]["checks"]))
        book = load_workbook(self.output / "口コミ分析.xlsx")
        self.assertNotIn("返信案", book.sheetnames)
        self.assertEqual(len(book.sheetnames), 6)
        self.assertEqual(book["分類済み口コミ"].max_row, 4)
        book.close()

    def test_reply_preference_is_saved_at_preparation_and_is_opt_in(self):
        default = self.root / "without-replies"
        enabled = self.root / "with-replies"
        with redirect_stdout(io.StringIO()):
            report_cli.report_main(["prepare", "--input", str(self.save_csv()), "--output-dir", str(default)])
            report_cli.report_main(["prepare", "--input", str(self.save_csv()), "--output-dir", str(enabled), "--with-replies"])
        self.assertFalse(r.read_json(default / "manifest.json")["reply_drafts"])
        self.assertTrue(r.read_json(enabled / "manifest.json")["reply_drafts"])

    def test_workbook_preserves_cause_checks_and_editable_action_status(self):
        self.synthesis["concerns"] = [{"title": "提供時間", "observation": "提供が遅いという記述", "hypothesis": "提供順の確認が必要", "alternative": "混雑の可能性", "verify": "提供時間を記録する", "evidence_ids": ["R0001"]}]
        self.synthesis["actions"] = [{"id": "A01", "title": "提供時間を確認", "detail": "記録から運営を確かめる", "owner": "担当案", "timeframe": "0〜7日", "priority": "高", "priority_reason": "本文に待ち時間の指摘", "effort": "低", "kpi": "提供時間", "baseline": "未測定", "target": "まず計測する", "evidence_ids": ["R0001"]}]
        r.write_json(self.output / "synthesis.json", self.synthesis)
        r.render(self.output)
        book = load_workbook(self.output / "口コミ分析.xlsx")
        self.assertEqual(book["分析と提案"]["D2"].value, "提供順の確認が必要")
        self.assertEqual(book["分析と提案"]["F2"].value, "提供時間を記録する")
        self.assertEqual(book["改善アクション"]["L2"].value, "R0001")
        validations = list(book["改善アクション"].data_validations.dataValidation)
        self.assertEqual(len(validations), 1)
        self.assertIn("進行中", validations[0].formula1)
        self.assertEqual(str(validations[0].sqref), "M2")
        self.assertFalse(any(cell.data_type == "f" for sheet in book for row in sheet for cell in row))
        book.close()

    def test_duplicate_source_ids_and_existing_manifest_are_not_overwritten(self):
        with self.assertRaisesRegex(ValueError, "上書き"):
            r.prepare(self.source, self.output)
        self.raw[1]["review_id"] = self.raw[0]["review_id"]
        self.save_source()
        with self.assertRaisesRegex(ValueError, "重複"):
            r.prepare(self.source, self.root / "other")

    def test_partial_batch_is_rejected_but_missing_batches_can_resume(self):
        (self.output / "annotations" / "batch-001.json").unlink()
        self.assertEqual(len(r.load_annotations(self.output, require_complete=False)[2]), 1)
        self.annotations.pop()
        self.save_annotations()
        with self.assertRaises(ValueError):
            r.load_annotations(self.output, require_complete=False)

    def save_csv(self, rows=None):
        path = self.root / "reviews.csv"
        with path.open("w", encoding="utf-8-sig", newline="") as stream:
            writer = csv.writer(stream)
            writer.writerow([title for _, title in cli.FIELDS])
            for row in rows if rows is not None else self.raw:
                writer.writerow([row.get(key, "") for key, _ in cli.FIELDS])
        return path

    def test_csv_preparation_preserves_original_and_never_collects(self):
        source = self.save_csv()
        output = self.root / "csv-analysis"
        with patch.object(cli, "run_collection", side_effect=AssertionError("no collection")), patch.object(report_cli.subprocess, "Popen", side_effect=AssertionError("no child AI")), redirect_stdout(io.StringIO()):
            self.assertEqual(report_cli.report_main(["prepare", "--input", str(source), "--output-dir", str(output)]), 0)
        self.assertEqual((output / "source.csv").read_bytes(), source.read_bytes())
        manifest, metadata, keys = r.workspace(output)
        self.assertEqual(len(keys), 3)
        self.assertEqual(keys["R0001"]["text"], self.raw[0]["text"])
        self.assertEqual(keys["R0001"]["author"], "=AUTHOR()")
        self.assertEqual(manifest["coverage"]["input_format"], "csv")
        self.assertIsNone(manifest["coverage"]["displayed_total"])
        self.assertFalse(manifest["coverage"]["full_coverage_verified"])
        (output / "source.csv").write_bytes(b"changed")
        with self.assertRaisesRegex(ValueError, "入力CSV"):
            r.workspace(output)

    def test_url_report_commands_are_rejected_before_collection(self):
        with patch.object(cli, "run_collection", side_effect=AssertionError("no collection")), redirect_stderr(io.StringIO()):
            for arguments in (["https://maps.app.goo.gl/store"], ["prepare", "https://maps.app.goo.gl/store"], ["resume", str(self.output)]):
                with self.subTest(arguments=arguments), self.assertRaises(SystemExit):
                    report_cli.report_main(arguments)

    def test_csv_missing_ids_keep_distinct_identical_rows(self):
        path = self.root / "minimal.csv"
        path.write_text('rating,text\n5,"良い\n料理"\n5,"良い\n料理"\n4,\n', encoding="utf-8")
        metadata, rows, _ = r.source_data(path)
        self.assertEqual(len(rows), 3)
        self.assertEqual(rows[0]["text"], "良い\n料理")
        self.assertEqual(len({row["review_id"] for row in rows}), 3)
        self.assertFalse(rows[2]["text_may_be_truncated"])

    def test_csv_invalid_rows_and_mixed_places_are_rejected(self):
        path = self.root / "invalid.csv"
        for text in ('rating,other\n5,hello\n', 'rating,text\n6,hello\n', 'rating,text\n5,hello,extra\n', 'rating,text,text\n5,a,b\n', 'rating,text,review_id\n5,a,same\n4,b,same\n', 'rating,text,place_name\n5,a,A\n4,b,B\n', 'rating,text\n'):
            with self.subTest(text=text):
                path.write_text(text, encoding="utf-8")
                with self.assertRaises(ValueError):
                    r.source_data(path)

    def test_csv_report_renders_every_row_in_one_workbook(self):
        source = self.save_csv()
        output = self.root / "csv-render"
        r.prepare(source, output, reply_drafts=False)
        manifest = r.read_json(output / "manifest.json")
        r.write_json(output / "annotations" / "batch-001.json", {"source_sha256": manifest["source_sha256"], "reviews": self.annotations})
        synthesis = dict(self.synthesis, source_sha256=manifest["source_sha256"])
        r.write_json(output / "synthesis.json", synthesis)
        html = r.render(output).read_text()
        self.assertIn("入力CSVの全", html)
        book = load_workbook(output / "口コミ分析.xlsx")
        self.assertEqual(len(book.sheetnames), 6)
        self.assertEqual(book["分類済み口コミ"].max_row, 4)
        self.assertEqual(book["分類済み口コミ"]["D2"].value, 5)
        book.close()

    def test_skill_install_updates_own_skill_and_refuses_foreign_skill(self):
        parent = self.root / "skills"
        target = report_cli.install_skill(parent)
        self.assertTrue((target / "references" / "schema.md").is_file())
        report_cli.install_skill(parent)
        (target / ".google-maps-reviews-skill").unlink()
        with self.assertRaisesRegex(ValueError, "上書き"):
            report_cli.install_skill(parent)



if __name__ == "__main__":
    unittest.main()
