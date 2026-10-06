import csv
import io
import json
import tempfile
import unittest
from contextlib import redirect_stderr, redirect_stdout
from pathlib import Path
from unittest.mock import MagicMock, patch

from openpyxl import load_workbook
from google_maps_reviews import cli


class ContentPolicyTest(unittest.TestCase):
    url = "https://www.google.com/maps/place/Test/data=!1s0x123:0x456"

    def row(self, key, text="", date="2026-09-15", reply=""):
        return {"review_id": key, "author": "架空", "rating": 4, "text": text,
                "date_text": date, "owner_reply": reply, "text_may_be_truncated": False}

    def browser_run(self, batches, options, total):
        page = MagicMock(url=self.url)
        page.evaluate.side_effect = [{"place_name": "テスト", "source_url": self.url,
                                      "displayed_total": total, "reviews": rows} for rows in batches]
        context = MagicMock(pages=[page])
        output, errors = io.StringIO(), io.StringIO()
        with tempfile.TemporaryDirectory() as folder, patch("playwright.sync_api.sync_playwright"), \
                patch.object(cli, "collect_from_local_service", return_value=False), \
                patch.object(cli, "collection_browser") as launch, patch.object(cli, "prepare_reviews"), \
                patch.object(cli, "expand_text"), patch.object(cli, "blocked", return_value=False), \
                patch.object(cli, "open_full_reviews", return_value=False), \
                patch.object(cli, "reviews_restricted", return_value=False), \
                patch.object(cli, "review_scroll_state", return_value={"at_end": False}), \
                patch.object(cli, "scroll_reviews"), patch.object(cli.time, "sleep"), \
                redirect_stdout(output), redirect_stderr(errors):
            launch.return_value.__enter__.return_value = context
            code = cli.main([self.url, *options, "--output-dir", folder])
            data = json.loads(next(Path(folder).glob("*.json")).read_text())
            groups = [data.get(key, []) for key in ("reviews", "uncertain_reviews", "excluded_reviews")]
            self.assertTrue(all(row["text"].strip() for group in groups for row in group))
            with next(Path(folder).glob("*.csv")).open(encoding="utf-8-sig") as stream:
                csv_rows = list(csv.DictReader(stream))
                self.assertTrue(all(row["口コミ本文"].strip() for row in csv_rows))
                self.assertEqual(len(csv_rows), data["metadata"]["count"])
            book = load_workbook(next(Path(folder).glob("*.xlsx")), read_only=True)
            excel_ids = []
            for sheet in book:
                if sheet.title == "取得情報":
                    continue
                for row in sheet.iter_rows(min_row=2, values_only=True):
                    self.assertTrue(row[4].strip())
                    excel_ids.append(row[6])
            book.close()
            self.assertEqual(set(excel_ids), {row["review_id"] for group in groups for row in group})
            self.assertEqual(len(excel_ids), len(set(excel_ids)))
            return code, data, output.getvalue(), errors.getvalue(), page.evaluate.call_count

    def test_all_and_visible_only_exclude_blank_and_owner_reply_only(self):
        rows = [self.row("first", "良い料理"), self.row("stars"), self.row("spaces", " \n\u3000"),
                self.row("owner-only", reply="ありがとうございます"), self.row("last", "提供が遅い")]
        for options, total in ((["--all"], 5), (["--visible-only"], 10)):
            with self.subTest(options=options):
                code, data, _, _, _ = self.browser_run([rows], options, total)
                self.assertEqual(code, 0)
                self.assertEqual([row["review_id"] for row in data["reviews"]], ["first", "last"])
                self.assertEqual(data["metadata"]["scanned_count"], 5)
                self.assertEqual(data["metadata"]["scanned_text_review_count"], 2)
                self.assertEqual(data["metadata"]["excluded_rating_only_count"], 3)
                self.assertEqual(data["metadata"]["rating_only_count"], 0)

    def test_max_counts_bodies_and_does_not_stop_at_first_rating_only_batch(self):
        stars = [self.row("star-1"), self.row("star-2")]
        one_body = stars + [self.row("text-1", "本文1"), self.row("star-3")]
        complete = one_body + [self.row("text-2", "本文2")]
        code, data, _, _, calls = self.browser_run([stars, one_body, complete], ["--max", "2"], 5)
        self.assertEqual((code, calls), (0, 3))
        self.assertEqual([row["review_id"] for row in data["reviews"]], ["text-1", "text-2"])
        self.assertEqual(data["metadata"]["excluded_rating_only_count"], 3)
        self.assertTrue(data["metadata"]["scan_coverage_verified"])

    def test_max_applies_to_saved_bodies_after_a_larger_dom_batch(self):
        rows = [self.row("star"), *[self.row(f"text-{i}", f"本文{i}") for i in range(3)]]
        code, data, _, _, _ = self.browser_run([rows], ["--max", "2"], 4)
        self.assertEqual(code, 0)
        self.assertEqual([row["review_id"] for row in data["reviews"]], ["text-0", "text-1"])
        self.assertEqual(data["metadata"]["scanned_text_review_count"], 3)
        self.assertTrue(data["metadata"]["scan_coverage_verified"])
        self.assertFalse(data["metadata"]["full_coverage_verified"])

    def test_max_succeeds_when_all_bodies_are_fewer_than_requested(self):
        rows = [self.row("star"), self.row("text", "本文")]
        code, data, _, _, _ = self.browser_run([rows], ["--max", "10"], 2)
        self.assertEqual(code, 0)
        self.assertEqual(len(data["reviews"]), 1)
        self.assertTrue(data["metadata"]["full_coverage_verified"])

    def test_max_retries_expansion_before_finishing_at_the_body_limit(self):
        collapsed = dict(self.row("text", "長い本文…"), text_may_be_truncated=True)
        expanded = self.row("text", "長い本文の全文")
        code, data, _, _, calls = self.browser_run([[collapsed], [expanded]], ["--max", "1"], 10)
        self.assertEqual((code, calls), (0, 2))
        self.assertEqual(data["reviews"][0]["text"], "長い本文の全文")
        self.assertEqual(data["metadata"]["saved_text_truncated_count"], 0)

    def test_unresolved_truncation_is_not_reported_as_collection_success(self):
        collapsed = dict(self.row("text", "長い本文…"), text_may_be_truncated=True)
        code, data, _, errors, calls = self.browser_run([[collapsed]] * 4, ["--all"], 1)
        self.assertEqual((code, calls), (2, 4))
        self.assertTrue(data["metadata"]["scan_coverage_verified"])
        self.assertEqual(data["metadata"]["saved_text_truncated_count"], 1)
        self.assertTrue(data["metadata"]["incomplete"])
        self.assertIn("本文の省略が残る", errors)

    def test_max_timeout_below_requested_count_is_partial_and_keeps_max_in_retry(self):
        rows = [self.row("star"), self.row("text", "本文")]
        with patch.object(cli.time, "monotonic", side_effect=[0, 10]):
            code, data, _, errors, _ = self.browser_run([rows], ["--max", "2", "--timeout", "1"], 5)
        self.assertEqual(code, 2)
        self.assertEqual(data["metadata"]["count"], 1)
        self.assertTrue(data["metadata"]["incomplete"])
        self.assertIn("--max 2", errors)
        self.assertNotIn("--all", errors)

    def test_rating_only_store_saves_an_empty_csv_and_reports_zero_bodies(self):
        rows = [self.row(str(i)) for i in range(3)]
        code, data, output, _, _ = self.browser_run([rows], ["--all"], 3)
        self.assertEqual(code, 0)
        self.assertEqual(data["reviews"], [])
        self.assertEqual(data["metadata"]["excluded_rating_only_count"], 3)
        self.assertTrue(data["metadata"]["full_coverage_verified"])
        self.assertIn("本文あり0件", output)

    def test_period_classifies_only_bodies_and_keeps_all_body_ids(self):
        rows = [self.row("in", "本文1"), self.row("boundary", "本文2", date="1 か月前"),
                self.row("unknown", "本文3", date="最終編集: 1 年前"),
                self.row("out", "本文4", date="2020-01-01"),
                self.row("stars-edited", date="最終編集: 1 年前"), self.row("stars-old", date="2020-01-01")]
        code, data, _, _, _ = self.browser_run([rows], ["--from", "2026-09-01", "--to", "2026-09-30"], 6)
        self.assertEqual(code, 2)
        self.assertEqual([row["review_id"] for row in data["reviews"]], ["in"])
        self.assertEqual({row["review_id"] for row in data["uncertain_reviews"]}, {"boundary", "unknown"})
        self.assertEqual([row["review_id"] for row in data["excluded_reviews"]], ["out"])
        self.assertEqual(data["metadata"]["scanned_count"], 6)
        self.assertEqual(data["metadata"]["scanned_text_review_count"], 4)
        self.assertEqual(data["metadata"]["period_unknown_date_count"], 1)
        self.assertEqual(data["metadata"]["excluded_rating_only_count"], 2)

    def test_edited_rating_only_date_does_not_make_exact_body_period_uncertain(self):
        rows = [self.row("text", "本文"), self.row("star", date="最終編集: 1 年前")]
        code, data, _, _, _ = self.browser_run([rows], ["--from", "2026-09-01", "--to", "2026-09-30"], 2)
        self.assertEqual(code, 0)
        self.assertTrue(data["metadata"]["period_date_accuracy_verified"])
        self.assertEqual(data["metadata"]["period_unknown_date_count"], 0)

    def test_service_export_uses_same_body_only_policy(self):
        rows = [self.row("text", "本文"), self.row("star"), self.row("owner-only", reply="返信")]
        event = {"type": "done", "data": {"place": "テスト", "sourceUrl": self.url,
                 "displayedTotal": 3, "reviews": rows, "reason": "画面の総件数と保存件数が一致", "verified": True}}
        opener = MagicMock()
        opener.open.side_effect = [io.BytesIO(b'{"ready":true,"version":"1.0.0","busy":false}'),
                                  io.BytesIO(b"data: " + json.dumps(event).encode() + b"\n\n")]
        with tempfile.TemporaryDirectory() as folder, patch.object(cli, "build_opener", return_value=opener), \
                patch.object(cli, "collection_browser", side_effect=AssertionError("browser not allowed")), \
                redirect_stdout(io.StringIO()):
            self.assertEqual(cli.main([self.url, "--all", "--output-dir", folder]), 0)
            data = json.loads(next(Path(folder).glob("*.json")).read_text())
            self.assertEqual([row["review_id"] for row in data["reviews"]], ["text"])
            self.assertEqual(data["metadata"]["excluded_rating_only_count"], 2)
            self.assertIn("読取件数", data["metadata"]["stop_reason"])

    def test_export_guard_rejects_blank_rows_in_every_body_only_group(self):
        for group in ("reviews", "period_uncertain_reviews", "period_excluded_reviews"):
            with self.subTest(group=group), tempfile.TemporaryDirectory() as folder:
                metadata = {"review_filter": "text_only"}
                rows = [self.row("text", "本文")]
                if group == "reviews":
                    rows.append(self.row("empty"))
                else:
                    metadata[group] = [self.row("empty")]
                with self.assertRaisesRegex(ValueError, "本文のない"):
                    cli.export_reviews(rows, metadata, Path(folder), "test")
                self.assertEqual(list(Path(folder).iterdir()), [])


if __name__ == "__main__":
    unittest.main()
