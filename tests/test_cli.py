import csv
import importlib.util
import io
import json
import subprocess
import sys
import tempfile
import unittest
import zipfile
from contextlib import redirect_stdout, redirect_stderr
from importlib.resources import files
from pathlib import Path
from unittest.mock import patch

from openpyxl import load_workbook
from google_maps_reviews import cli, console


class InstalledCliTest(unittest.TestCase):
    def test_extractor_is_available_in_installed_package(self):
        source = files("google_maps_reviews").joinpath("extract_reviews.js").read_text(encoding="utf-8")
        self.assertIn("data-review-id", source)
        self.assertIn("displayed_total", source)

    def test_command_runs_from_another_directory_and_exports(self):
        with tempfile.TemporaryDirectory() as folder:
            result = subprocess.run(
                ["google-maps-reviews", "--demo", "--output-dir", folder],
                cwd=folder, capture_output=True, text=True,
            )
            self.assertEqual(result.returncode, 0, result.stderr)
            paths = list(Path(folder).glob("*.json"))
            self.assertEqual(len(paths), 1)
            data = json.loads(paths[0].read_text(encoding="utf-8"))
            self.assertTrue(data["metadata"]["sample"])
            self.assertEqual(data["metadata"]["count"], 1)
            self.assertEqual(len(data["reviews"]), 1)
            with paths[0].with_suffix(".csv").open(encoding="utf-8-sig", newline="") as stream:
                self.assertEqual(len(list(csv.reader(stream))), 2)
            book = load_workbook(paths[0].with_suffix(".xlsx"))
            self.assertEqual(book["口コミ"].max_row, 2)
            book.close()

    def test_conflicting_scope_options_are_rejected(self):
        for options in [("--all", "--max", "5"), ("--all", "--visible-only")]:
            result = subprocess.run(
                ["google-maps-reviews", *options], capture_output=True, text=True,
            )
            self.assertEqual(result.returncode, 2)

    def test_non_terminal_never_waits_for_interactive_input(self):
        for options, expected in [([], 0), (["--help"], 0), (["--interactive"], 2),
                                  (["--manual", "https://www.google.com/maps/test"], 2),
                                  (["--max", "5"], 2)]:
            result = subprocess.run(["google-maps-reviews", *options], stdin=subprocess.DEVNULL,
                                    capture_output=True, text=True, timeout=5)
            self.assertEqual(result.returncode, expected, result.stderr)
        result = subprocess.run([sys.executable, "-m", "google_maps_reviews", "--help"],
                                stdin=subprocess.DEVNULL, capture_output=True, text=True, timeout=5)
        self.assertEqual(result.returncode, 0, result.stderr)

    def test_interactive_collection_saves_and_reuses_all_preferences(self):
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "settings.json"
            output = Path(folder) / "reviews"
            responses = ["1", "https://www.google.com/maps/test", "2", "7", "45", "1.5", str(output), "1", "2", "4"]
            with patch.object(sys.stdin, "isatty", return_value=True), patch("builtins.input", side_effect=responses), redirect_stdout(io.StringIO()):
                self.assertEqual(cli.main(["--interactive", "--demo", "--config", str(path)]), 0)
            settings = console.load_settings(path)
            self.assertEqual(settings, {"url": "https://www.google.com/maps/test", "scope": "max", "max": 7,
                                       "timeout": 45, "delay": 1.5, "output_dir": str(output), "browser": "chrome", "manual": True})
            self.assertEqual(path.stat().st_mode & 0o777, 0o600)
            with patch.object(sys.stdin, "isatty", return_value=False), redirect_stdout(io.StringIO()):
                self.assertEqual(cli.main(["--use-settings", "--demo", "--max", "3", "--config", str(path)]), 0)
            maxima = {json.loads(file.read_text(encoding="utf-8"))["metadata"]["requested_max"] for file in output.glob("*.json")}
            self.assertEqual(maxima, {3, 7})

    def test_saved_settings_are_opt_in_and_explicit_flags_override_scope(self):
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "settings.json"
            settings = dict(console.default_settings(), scope="all", max=50, timeout=40, delay=3.0, browser="chromium")
            console.save_settings(path, settings)
            with patch.object(sys.stdin, "isatty", return_value=False), patch.object(cli, "run_collection", return_value=0) as collecting:
                self.assertEqual(cli.main(["https://www.google.com/maps/test", "--config", str(path)]), 0)
                args = collecting.call_args.args[0]
                self.assertEqual((args.max, args.timeout, args.browser, args.all), (100, 300, "chrome", False))
                self.assertEqual(cli.main(["https://www.google.com/maps/test", "--use-settings", "--max=12", "--timeout", "99", "--config", str(path)]), 0)
                args = collecting.call_args.args[0]
                self.assertEqual((args.max, args.timeout, args.browser, args.all), (12, 99, "chromium", False))
                self.assertEqual(cli.main(["https://www.google.com/maps/test", "--demo", "--use-settings", "--time", "99", "--bro=chrome", "--man", "--config", str(path)]), 0)
                args = collecting.call_args.args[0]
                self.assertEqual((args.timeout, args.browser, args.manual), (99, "chrome", True))
            with patch.object(sys.stdin, "isatty", return_value=True), patch.object(console, "edit_settings", side_effect=lambda current: current), patch.object(cli, "run_collection", return_value=0) as collecting, patch("builtins.input", side_effect=["1", "4"]), redirect_stdout(io.StringIO()):
                self.assertEqual(cli.main(["--interactive", "--time", "99", "--bro", "chrome", "--man", "--config", str(path)]), 0)
                args = collecting.call_args.args[0]
                self.assertEqual((args.timeout, args.browser, args.manual), (99, "chrome", True))

    def test_menu_exit_eof_and_bad_configuration_do_not_start_a_browser(self):
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "settings.json"
            with patch.dict("os.environ", {"GOOGLE_MAPS_REVIEWS_CONFIG": str(path)}), patch.object(sys.stdin, "isatty", return_value=True), patch.object(cli, "run_collection") as collecting, redirect_stdout(io.StringIO()):
                with patch("builtins.input", return_value="4"):
                    self.assertEqual(cli.main([]), 0)
                with patch("builtins.input", side_effect=EOFError):
                    self.assertEqual(cli.main(["--interactive", "--config", str(path)]), 0)
                collecting.assert_not_called()
            path.write_text('DUMMY_REVIEW_SECRET', encoding="utf-8")
            errors = io.StringIO()
            with redirect_stderr(errors):
                self.assertEqual(cli.main(["settings", "show", "--config", str(path)]), 1)
            self.assertNotIn("DUMMY_REVIEW_SECRET", errors.getvalue())
            with redirect_stdout(io.StringIO()):
                self.assertEqual(cli.main(["settings", "reset", "--config", str(path)]), 0)
            self.assertEqual(console.load_settings(path), console.default_settings())

    def test_settings_edit_requires_terminal_and_setup_uses_current_python(self):
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "settings.json"
            with patch.object(sys.stdin, "isatty", return_value=False), redirect_stderr(io.StringIO()):
                with self.assertRaises(SystemExit) as result:
                    cli.main(["settings", "edit", "--config", str(path)])
                self.assertEqual(result.exception.code, 2)
            with patch.object(console.subprocess, "run", return_value=subprocess.CompletedProcess([], 0)) as installing, redirect_stdout(io.StringIO()):
                self.assertEqual(cli.main(["setup", "--browser", "chromium", "--config", str(path)]), 0)
                installing.assert_called_once_with([sys.executable, "-m", "playwright", "install", "chromium"], check=False)
            self.assertEqual(console.load_settings(path)["browser"], "chromium")


class InstallerTest(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        script = Path(__file__).resolve().parents[1] / "install-cli.py"
        spec = importlib.util.spec_from_file_location("review_cli_installer", script)
        cls.installer = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(cls.installer)

    def test_unrelated_command_and_virtual_environment_are_not_overwritten(self):
        with tempfile.TemporaryDirectory() as folder:
            prefix = Path(folder)
            command = prefix / "bin/google-maps-reviews"
            command.parent.mkdir()
            command.write_text("another application's command", encoding="utf-8")
            with self.assertRaises(RuntimeError):
                self.installer.check_destination(prefix)
            self.assertEqual(command.read_text(encoding="utf-8"), "another application's command")
            command.unlink()
            environment = prefix / "share/google-maps-reviews/venv"
            environment.mkdir(parents=True)
            with self.assertRaises(RuntimeError):
                self.installer.check_destination(prefix)

    def test_prefix_isolates_command_virtual_environment_and_configuration(self):
        with tempfile.TemporaryDirectory(prefix="cli install space ") as folder:
            prefix = (Path(folder) / "user prefix").resolve()
            source = Path(__file__).resolve().parents[1]
            with patch.object(self.installer.venv, "EnvBuilder") as builder, patch.object(self.installer.subprocess, "run") as running, redirect_stdout(io.StringIO()):
                self.assertEqual(self.installer.main(["--prefix", str(prefix), "--source", str(source), "--browser", "chrome"]), 0)
                self.assertEqual(self.installer.main(["--prefix", str(prefix), "--source", str(source), "--browser", "chrome"]), 0)
            environment = prefix / "share/google-maps-reviews/venv"
            settings = prefix / "share/google-maps-reviews/settings.json"
            command = prefix / "bin/google-maps-reviews"
            self.assertEqual(command.read_text(encoding="utf-8"), self.installer.launcher_text(environment / "bin/google-maps-reviews", settings))
            self.assertEqual(command.stat().st_mode & 0o777, 0o755)
            builder.return_value.create.assert_called_with(environment)
            self.assertEqual(running.call_args.args[0], [str(environment / "bin/google-maps-reviews"), "setup", "--browser", "chrome", "--config", str(settings)])

    def test_authenticated_gh_archive_is_extracted_without_git_checkout_changes(self):
        with tempfile.TemporaryDirectory() as folder:
            payload = io.BytesIO()
            with zipfile.ZipFile(payload, "w") as archive:
                archive.writestr("source-root/pyproject.toml", "[project]\nname='google-maps-reviews'\n")
                archive.writestr("source-root/src/google_maps_reviews/cli.py", "# source\n")
            def download(command, stdout, check):
                self.assertEqual(command, ["gh", "api", "repos/atsuki-sakai/google-maps-reviews/zipball/main"])
                self.assertTrue(check)
                stdout.write(payload.getvalue())
            with patch.object(self.installer.subprocess, "run", side_effect=download):
                source = self.installer.fetch_source("main", Path(folder))
            self.assertTrue((source / "src/google_maps_reviews/cli.py").is_file())

    def test_archive_cannot_write_outside_temporary_source_folder(self):
        with tempfile.TemporaryDirectory() as folder:
            payload = io.BytesIO()
            with zipfile.ZipFile(payload, "w") as archive:
                archive.writestr("../outside.txt", "unexpected")
            def download(_command, stdout, check):
                stdout.write(payload.getvalue())
            with patch.object(self.installer.subprocess, "run", side_effect=download), self.assertRaises(RuntimeError):
                self.installer.fetch_source("main", Path(folder))
            self.assertFalse((Path(folder) / "outside.txt").exists())


if __name__ == "__main__":
    unittest.main()
