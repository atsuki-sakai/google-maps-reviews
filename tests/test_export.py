import argparse
import csv
import json
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook
from google_maps_reviews.cli import export_reviews, maps_url, merge_reviews, next_stagnant_count, full_coverage_verified


class ReviewsTest(unittest.TestCase):
    def test_total_must_match_unique_review_ids(self):
        rows = [{"review_id": str(i)} for i in range(331)]
        self.assertTrue(full_coverage_verified(rows, 331))
        self.assertFalse(full_coverage_verified(rows[:100], 331))
        self.assertFalse(full_coverage_verified(rows, None))
        self.assertFalse(full_coverage_verified(rows[:-1] + [rows[0]], 331))
        self.assertFalse(full_coverage_verified(rows[:-1] + [{"review_id": ""}], 331))

    def test_long_loaded_reviews_do_not_trigger_early_stop(self):
        stalled = 0
        # Twenty passes through long loaded reviews must not be mistaken for the end.
        for _ in range(20):
            stalled = next_stagnant_count(stalled, added=0, state={"at_end": False})
            self.assertEqual(stalled, 0)
        for expected in range(1, 9):
            stalled = next_stagnant_count(stalled, added=0, state={"at_end": True})
            self.assertEqual(stalled, expected)
        self.assertEqual(next_stagnant_count(stalled, added=5, state={"at_end": True}), 0)

    def test_expanded_review_and_rating_only_are_preserved(self):
        rows = {}
        full = {"review_id": "same", "author": "投稿者", "rating": 5,
                "text": "長い本文\n全文を保存", "owner_reply": "ありがとうございます", "text_may_be_truncated": False}
        self.assertEqual(merge_reviews(rows, [full], "店舗", "URL"), 1)
        collapsed = dict(full, text="長い本文", owner_reply="", text_may_be_truncated=True)
        self.assertEqual(merge_reviews(rows, [collapsed], "店舗", "URL"), 0)
        self.assertEqual(rows["same"]["text"], full["text"])
        self.assertFalse(rows["same"]["text_may_be_truncated"])
        self.assertEqual(rows["same"]["owner_reply"], full["owner_reply"])
        self.assertEqual(merge_reviews(rows, [{"review_id": "rating", "author": "星のみ", "rating": 3, "text": ""}], "店舗", "URL"), 1)

    def test_export_roundtrip_and_formula_text(self):
        body = '=HYPERLINK("https://example.com")\n口コミ,日本語 😀\x01' + "あ" * 33000
        rows = [{"place_name": "架空店舗", "author": "投稿者", "rating": 4,
                 "text": body, "owner_reply": "返信\n2行目", "text_may_be_truncated": False}]
        with tempfile.TemporaryDirectory() as folder:
            files = export_reviews(rows, {"count": 1, "stop_reason": "テスト", "sample": True}, Path(folder), "test")
            self.assertTrue(all(path.exists() for path in files))
            with files[0].open(encoding="utf-8-sig", newline="") as file:
                csv_rows = list(csv.reader(file))
            self.assertEqual(csv_rows[1][4], "'" + body)
            original = json.loads(files[2].read_text(encoding="utf-8"))
            self.assertEqual(original["reviews"][0]["text"], body)
            book = load_workbook(files[1])
            self.assertEqual(book.sheetnames, ["口コミ", "取得情報"])
            self.assertEqual(book["口コミ"]["C2"].value, 4)
            self.assertEqual(book["口コミ"]["F2"].value, "返信\n2行目")
            self.assertEqual(book["口コミ"]["E2"].data_type, "s")
            self.assertEqual(len(book["口コミ"]["E2"].value), 32767)
            self.assertNotIn("\x01", book["口コミ"]["E2"].value)
            self.assertFalse(any(cell.data_type == "f" for sheet in book for row in sheet for cell in row))

    def test_google_maps_urls_only(self):
        self.assertEqual(maps_url("  https://maps.app.goo.gl/example\n"), "https://maps.app.goo.gl/example")
        for url in ["https://maps.app.goo.gl/example", "https://www.google.com/maps/place/example", "https://maps.google.co.jp/?q=example"]:
            self.assertEqual(maps_url(url), url)
        for url in ["https://example.com/maps", "https://google.com.example.com/maps", "file:///tmp/a", "https://www.google.com/search?q=example"]:
            with self.assertRaises(argparse.ArgumentTypeError):
                maps_url(url)


if __name__ == "__main__":
    unittest.main()
