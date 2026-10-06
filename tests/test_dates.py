import csv
import argparse
import io
import json
import tempfile
import unittest
from contextlib import redirect_stdout, redirect_stderr
from datetime import date
from pathlib import Path
from unittest.mock import MagicMock, patch

from openpyxl import load_workbook
from google_maps_reviews import cli, console
from google_maps_reviews.dates import date_bounds, iso_date, partition_period, select_period


class DateRangeTest(unittest.TestCase):
    def setUp(self):
        network = patch.object(cli, "build_opener", side_effect=OSError("単体テストでは収集サービスに接続しません"))
        network.start()
        self.addCleanup(network.stop)

    def test_relative_labels_remain_ranges_and_edited_dates_are_unknown(self):
        today = date(2026, 10, 6)
        for text in ("5 か月前", "a month ago", "2 years ago", "3 日前", "1週間前"):
            with self.subTest(text=text):
                lower, upper, precision = date_bounds(text, today)
                self.assertLess(lower, upper)
                self.assertIn("推定", precision)
                self.assertLessEqual(upper, today)
        self.assertEqual(date_bounds("最終編集: 1 年前", today)[:2], (None, None))
        self.assertEqual(date_bounds("not a date", today)[:2], (None, None))

    def test_exact_date_and_calendar_month_end(self):
        self.assertEqual(date_bounds("2026年9月30日", date(2026, 10, 6))[:2], (date(2026, 9, 30),) * 2)
        lower, upper, _ = date_bounds("1か月前", date(2024, 3, 31))
        self.assertLessEqual(lower, date(2024, 1, 31))
        self.assertEqual(upper, date(2024, 3, 17))
        with self.assertRaises(argparse.ArgumentTypeError):
            iso_date("2026-02-30")
        with self.assertRaises(argparse.ArgumentTypeError):
            iso_date("2026-1-1")

    def test_one_year_ago_is_inside_the_users_multi_year_period(self):
        anchor = date(2026, 10, 6)
        rows = [{"review_id": str(i), "date_text": text} for i, text in enumerate(
            ("5 か月前", "1 年前", "1 年前", "1 年前", "5 か月前"))]
        selected, uncertain, excluded = select_period(rows, "2023-01-01", "2026-09-30", anchor)
        self.assertEqual(len(selected), 5)
        self.assertEqual((uncertain, excluded), ([], 0))
        old_year = selected[1]
        self.assertEqual(old_year["date_earliest"], "2024-10-05")
        self.assertEqual(old_year["date_latest"], "2026-04-08")
        self.assertIn("丸め仮定", old_year["date_precision"])
        # A narrower boundary remains uncertain instead of fabricating a date.
        selected, uncertain, _ = select_period(rows[1:2], "2025-01-01", "2025-12-31", anchor)
        self.assertEqual((len(selected), len(uncertain)), (0, 1))

    def test_hour_labels_do_not_all_collapse_to_today(self):
        lower, upper, _ = date_bounds("96 hours ago", date(2026, 10, 6))
        self.assertLessEqual(lower, date(2026, 10, 2))
        self.assertLess(upper, date(2026, 10, 6))

    def test_period_includes_endpoints_and_retains_boundary_unknown(self):
        rows = [{"review_id": str(i), "date_text": text} for i, text in enumerate(
            ("2026-09-01", "2026-09-30", "2026-08-31", "1 か月前", "最終編集: 1 年前"))]
        selected, uncertain, excluded = select_period(rows, "2026-09-01", "2026-09-30", date(2026, 10, 6))
        self.assertEqual([row["review_id"] for row in selected], ["0", "1"])
        self.assertEqual([row["review_id"] for row in uncertain], ["3", "4"])
        self.assertEqual(excluded, 1)
        self.assertNotIn("date_precision", rows[0])

    def test_period_conflicts_and_reversed_range_fail_before_browser(self):
        with patch.object(cli, "collection_browser", side_effect=AssertionError("no browser")), redirect_stderr(io.StringIO()):
            for arguments in (["--from", "2026-01-01", "--all"], ["--from", "2026-01-01", "--max", "5"], ["--to", "2026-01-01", "--visible-only"], ["--from", "2026-09-30", "--to", "2026-09-01"]):
                with self.subTest(arguments=arguments), self.assertRaises(SystemExit):
                    cli.main(["https://maps.app.goo.gl/store"] + arguments)

    def test_settings_do_not_leak_period_into_an_explicit_all_run(self):
        settings = dict(console.default_settings(), scope="period", date_from="2026-01-01", date_to="2026-06-30")
        args = console.apply_settings(cli.build_parser().parse_args(["--all"]), settings, ["--all"])
        self.assertIsNone(args.date_from)
        self.assertTrue(args.all)
        args = console.apply_settings(cli.build_parser().parse_args([]), settings, [])
        self.assertEqual(args.date_from, "2026-01-01")
        self.assertFalse(args.all)

    def test_interactive_period_settings_can_be_saved(self):
        answers = ["https://maps.app.goo.gl/store", "4", "2026-01-01", "2026-09-30", "", "", "", "", ""]
        with patch("builtins.input", side_effect=answers), redirect_stdout(io.StringIO()):
            settings = console.edit_settings(console.default_settings())
        self.assertEqual(settings["scope"], "period")
        self.assertEqual(settings["date_to"], "2026-09-30")

    def test_period_collection_scans_all_and_saves_uncertainty_in_same_workbook(self):
        page = MagicMock(url="https://www.google.com/maps/test")
        reviews = [{"review_id": str(i), "author": "架空", "rating": 4, "text": "架空の口コミ本文", "date_text": text} for i, text in enumerate(
            ("2026-09-01", "2026-09-30", "2026-08-31", "最終編集: 1 年前"))]
        reviews[2]["text"] = '=HYPERLINK("https://example.com")'
        reviews.append({"review_id": "rating-only", "author": "架空", "rating": 5, "text": "", "date_text": "最終編集: 1 年前"})
        page.evaluate.return_value = {"place_name": "テスト", "source_url": page.url, "displayed_total": 5, "reviews": reviews}
        context = MagicMock(pages=[page])
        with tempfile.TemporaryDirectory() as folder, patch("playwright.sync_api.sync_playwright"), patch.object(cli, "collection_browser") as launch, patch.object(cli, "prepare_reviews"), patch.object(cli, "expand_text"), patch.object(cli, "blocked", return_value=False), redirect_stdout(io.StringIO()) as output, redirect_stderr(io.StringIO()):
            launch.return_value.__enter__.return_value = context
            code = cli.main([page.url, "--from", "2026-09-01", "--to", "2026-09-30", "--output-dir", folder])
            self.assertEqual(code, 2)
            data = json.loads(next(Path(folder).glob("*.json")).read_text())
            self.assertEqual(len(data["reviews"]), 2)
            self.assertEqual(len(data["uncertain_reviews"]), 1)
            self.assertEqual([row["review_id"] for row in data["excluded_reviews"]], ["2"])
            self.assertEqual(data["excluded_reviews"][0]["text"], reviews[2]["text"])
            exported = data["reviews"] + data["uncertain_reviews"] + data["excluded_reviews"]
            self.assertEqual(len(exported), 4)
            self.assertEqual({row["review_id"] for row in exported}, {"0", "1", "2", "3"})
            self.assertEqual(data["metadata"]["scanned_count"], 5)
            self.assertEqual(data["metadata"]["scanned_text_review_count"], 4)
            self.assertEqual(data["metadata"]["excluded_rating_only_count"], 1)
            self.assertEqual(data["metadata"]["missing_count"], 0)
            self.assertTrue(data["metadata"]["scan_coverage_verified"])
            self.assertFalse(data["metadata"]["full_coverage_verified"])
            self.assertFalse(data["metadata"]["period_date_accuracy_verified"])
            self.assertNotIn("保存件数が一致", output.getvalue())
            self.assertIn("一覧読取: 5件", output.getvalue())
            self.assertIn("期間外: 1件", output.getvalue())
            with next(Path(folder).glob("*.csv")).open(encoding="utf-8-sig", newline="") as stream:
                self.assertEqual(len(list(csv.DictReader(stream))), 2)
            book = load_workbook(next(Path(folder).glob("*.xlsx")))
            self.assertEqual(book.sheetnames, ["取得情報", "口コミ", "期間境界・日付不明", "期間外"])
            self.assertEqual(book["期間外"]["G2"].value, "2")
            self.assertEqual(book["期間外"]["E2"].value, reviews[2]["text"])
            self.assertEqual(book["期間外"]["E2"].data_type, "s")
            book.close()

    def test_excluded_rows_keep_dates_and_period_reason_without_mutating_input(self):
        original = [{"review_id": "old", "date_text": "2020-01-01"},
                    {"review_id": "recent", "date_text": "2026-10-01"}]
        selected, uncertain, excluded = partition_period(original, "2023-01-01", "2026-09-30", date(2026, 10, 6))
        self.assertEqual((selected, uncertain), ([], []))
        self.assertEqual([row["review_id"] for row in excluded], ["old", "recent"])
        self.assertEqual(excluded[1]["date_earliest"], "2026-10-01")
        self.assertIn("期間外", excluded[0]["period_match"])
        self.assertNotIn("period_match", original[0])

    def test_relative_dates_without_boundary_rows_still_do_not_claim_exact_period_success(self):
        page, context = self.restriction_fixture([1], expected=1)
        with tempfile.TemporaryDirectory() as folder, patch("playwright.sync_api.sync_playwright"), patch.object(cli, "collection_browser") as launch, patch.object(cli, "prepare_reviews"), patch.object(cli, "expand_text"), patch.object(cli, "blocked", return_value=False), redirect_stdout(io.StringIO()) as output, redirect_stderr(io.StringIO()) as errors:
            launch.return_value.__enter__.return_value = context
            code = cli.main([page.url, "--from", "2023-01-01", "--to", "2026-09-30", "--output-dir", folder])
            metadata = json.loads(next(Path(folder).glob("*.json")).read_text())["metadata"]
            self.assertEqual(code, 2)
            self.assertTrue(metadata["scan_coverage_verified"])
            self.assertTrue(metadata["period_selection_verified"])
            self.assertFalse(metadata["period_date_accuracy_verified"])
            self.assertEqual(metadata["period_uncertain_count"], 0)
            self.assertEqual(metadata["period_estimated_date_count"], 1)
            self.assertIn("期間内と推定: 1件", output.getvalue())
            self.assertIn("正確な投稿日による期間抽出は未確認", errors.getvalue())

    def test_empty_period_result_still_saves_valid_files_after_scan(self):
        selected, uncertain, excluded = select_period([{"date_text": "2020-01-01"}], "2026-01-01", "2026-09-30", date(2026, 10, 6))
        self.assertEqual((selected, uncertain, excluded), ([], [], 1))
        with tempfile.TemporaryDirectory() as folder:
            paths = cli.export_reviews([], {"count": 0}, Path(folder), "empty")
            self.assertEqual(len(paths), 3)
            with paths[0].open(encoding="utf-8-sig") as stream:
                self.assertEqual(list(csv.DictReader(stream)), [])

    def test_preview_opens_complete_list_and_reloads_stale_dom_once(self):
        page = MagicMock()
        button, sort = MagicMock(), MagicMock()
        page.get_by_role.side_effect = lambda role, name: MagicMock(first=sort if "並べ替え" in name.pattern else button)
        button.is_visible.return_value = True
        button.wait_for.side_effect = [TimeoutError("preview remains"), None]
        self.assertTrue(cli.open_full_reviews(page))
        page.reload.assert_called_once()
        button.click.assert_called_once()
        button.is_visible.return_value = False
        self.assertFalse(cli.open_full_reviews(page))

    def test_google_restriction_is_detected_even_when_review_cards_exist(self):
        page = MagicMock()
        page.locator.return_value.inner_text.return_value = "本文ありの口コミ\nGoogle マップの表示が制限されています。"
        self.assertTrue(cli.reviews_restricted(page))
        page.locator.return_value.inner_text.return_value = "口コミの本文のみ"
        self.assertFalse(cli.reviews_restricted(page))

    def restriction_fixture(self, counts, expected=10):
        page = MagicMock(url="https://www.google.com/maps/place/Test/data=!1s0x123:0x456")
        reviews = [{"review_id": str(i), "author": "架空", "rating": 4, "text": "架空の口コミ本文", "date_text": "1 年前"} for i in range(expected)]
        page.evaluate.side_effect = [{"place_name": "テスト", "source_url": page.url,
                                      "displayed_total": expected, "reviews": reviews[:count]} for count in counts]
        return page, MagicMock(pages=[page])

    def test_banner_does_not_interrupt_loading_and_full_list_is_opened_only_once(self):
        page, context = self.restriction_fixture([5, 5, 10, 15], expected=15)
        with tempfile.TemporaryDirectory() as folder, patch("playwright.sync_api.sync_playwright"), patch.object(cli, "collection_browser") as launch, patch.object(cli, "prepare_reviews"), patch.object(cli, "expand_text"), patch.object(cli, "blocked", return_value=False), patch.object(cli, "reviews_restricted", return_value=True), patch.object(cli, "recover_restricted_reviews", side_effect=AssertionError("読み込み中に操作を要求しました")) as restoring, patch.object(cli, "open_full_reviews", side_effect=[True, AssertionError("全口コミを再度開きました")]) as opening, patch.object(cli, "review_scroll_state", return_value={"at_end": False}), patch.object(cli, "scroll_reviews"), patch.object(cli.time, "sleep"), redirect_stdout(io.StringIO()):
            launch.return_value.__enter__.return_value = context
            self.assertEqual(cli.main([page.url, "--all", "--output-dir", folder]), 0)
            data = json.loads(next(Path(folder).glob("*.json")).read_text())
            self.assertEqual(len(data["reviews"]), 15)
            self.assertTrue(data["metadata"]["full_coverage_verified"])
            opening.assert_called_once()
            restoring.assert_not_called()

    def test_login_resume_keeps_collected_rows_and_excludes_manual_wait_from_timeout(self):
        page, context = self.restriction_fixture([5, 5, 5, 5, 8, 10])
        clock = [100.0]
        def recover(_page):
            clock[0] += 600  # Longer than the two-second collection timeout.
            return True
        with tempfile.TemporaryDirectory() as folder, patch("playwright.sync_api.sync_playwright"), patch.object(cli, "collection_browser") as launch, patch.object(cli, "prepare_reviews"), patch.object(cli, "expand_text"), patch.object(cli, "blocked", return_value=False), patch.object(cli, "reviews_restricted", side_effect=[True, False]), patch.object(cli, "recover_restricted_reviews", side_effect=recover) as restoring, patch.object(cli, "open_full_reviews", return_value=False), patch.object(cli, "review_scroll_state", return_value={"at_end": False}), patch.object(cli, "scroll_reviews"), patch.object(cli.time, "monotonic", side_effect=lambda: clock[0]), patch.object(cli.time, "sleep"), redirect_stdout(io.StringIO()), redirect_stderr(io.StringIO()):
            launch.return_value.__enter__.return_value = context
            code = cli.main([page.url, "--all", "--timeout", "2", "--output-dir", folder])
            data = json.loads(next(Path(folder).glob("*.json")).read_text())
            self.assertEqual(code, 0)
            self.assertEqual(len(data["reviews"]), 10)
            self.assertTrue(data["metadata"]["full_coverage_verified"])
            self.assertTrue(data["metadata"]["manual_recovery_attempted"])
            restoring.assert_called_once()

    def test_restriction_after_resume_saves_partial_once_and_returns_two(self):
        page, context = self.restriction_fixture([5] * 8)
        with tempfile.TemporaryDirectory() as folder, patch("playwright.sync_api.sync_playwright"), patch.object(cli, "collection_browser") as launch, patch.object(cli, "prepare_reviews"), patch.object(cli, "expand_text"), patch.object(cli, "blocked", return_value=False), patch.object(cli, "open_full_reviews", return_value=False), patch.object(cli, "review_scroll_state", return_value={"at_end": False}), patch.object(cli, "scroll_reviews"), patch.object(cli.time, "sleep"), patch.object(cli, "reviews_restricted", return_value=True), patch.object(cli, "recover_restricted_reviews", return_value=True) as restoring, redirect_stdout(io.StringIO()), redirect_stderr(io.StringIO()):
            launch.return_value.__enter__.return_value = context
            self.assertEqual(cli.main([page.url, "--all", "--output-dir", folder]), 2)
            data = json.loads(next(Path(folder).glob("*.json")).read_text())
            self.assertEqual(len(data["reviews"]), 5)
            self.assertEqual(data["metadata"]["missing_count"], 5)
            self.assertIn("表示制限が残", data["metadata"]["error"])
            restoring.assert_called_once()

    def test_nonterminal_restriction_never_prompts(self):
        with patch.object(cli.sys.stdin, "isatty", return_value=False), patch("builtins.input", side_effect=AssertionError("no input")):
            self.assertFalse(cli.recover_restricted_reviews(MagicMock()))

    def test_login_in_another_tab_refreshes_restricted_maps_once(self):
        page = MagicMock()
        with patch.object(cli.sys.stdin, "isatty", return_value=True), patch.object(cli, "prepare_reviews") as waiting, patch.object(cli, "reviews_restricted", return_value=True), redirect_stdout(io.StringIO()):
            self.assertTrue(cli.recover_restricted_reviews(page))
            waiting.assert_called_once_with(page, manual=True)
            page.reload.assert_called_once()

    def test_switching_places_during_login_never_merges_the_other_store(self):
        page, context = self.restriction_fixture([5] * 4)
        def recover(_page):
            page.url = "https://www.google.com/maps/place/Other/data=!1s0x789:0xabc"
            return True
        with tempfile.TemporaryDirectory() as folder, patch("playwright.sync_api.sync_playwright"), patch.object(cli, "collection_browser") as launch, patch.object(cli, "prepare_reviews"), patch.object(cli, "expand_text"), patch.object(cli, "blocked", return_value=False), patch.object(cli, "open_full_reviews", return_value=False), patch.object(cli, "review_scroll_state", return_value={"at_end": False}), patch.object(cli, "scroll_reviews"), patch.object(cli.time, "sleep"), patch.object(cli, "reviews_restricted", return_value=True), patch.object(cli, "recover_restricted_reviews", side_effect=recover), redirect_stdout(io.StringIO()), redirect_stderr(io.StringIO()):
            launch.return_value.__enter__.return_value = context
            self.assertEqual(cli.main([page.url, "--all", "--output-dir", folder]), 2)
            data = json.loads(next(Path(folder).glob("*.json")).read_text())
            self.assertIn("別施設", data["metadata"]["error"])
            self.assertEqual(page.evaluate.call_count, 4)


if __name__ == "__main__":
    unittest.main()
